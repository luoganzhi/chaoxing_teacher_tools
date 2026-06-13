#!/usr/bin/env python3
"""Local automation helper for the SCUJCC Chaoxing/Fanya portal.

This tool intentionally does not store usernames or passwords. The default
`login` command calls Chaoxing/Fanya HTTP login APIs, saves cookies to a local
cookie file, then later commands can reuse those cookies directly. Browser login
is kept as a fallback for captcha, SMS, QR, or two-factor flows.
"""

from __future__ import annotations

import argparse
import base64
import getpass
import importlib.util
import json
import os
import random
import re
import shutil
import subprocess
import sys
import time
from html import unescape
from urllib.parse import parse_qs, quote, unquote, urlencode, urljoin, urlparse


PORTAL_URL = "https://scujcc.fanya.chaoxing.com/portal"
DEFAULT_SESSION_NAME = "scujcc-chaoxing-portal"
LOGIN_PAGE_URL = "https://passport2.chaoxing.com/login"
FANYA_LOGIN_URL = "https://passport2.chaoxing.com/fanyalogin"
AUTH_CHECK_URL = "https://i.chaoxing.com/"
BACKCLAZZ_URL = "https://mooc1-api.chaoxing.com/mycourse/backclazzdata?view=json&rss=1"
COURSE_ENTRY_URL = "https://mooc1.chaoxing.com/course/isNewCourse"
MOOC2_BASE_URL = "https://mooc2-ans.chaoxing.com"
TRANSFER_KEY = "u2oh6Vu^HWe4_AES"
DEFAULT_GRADE_PLAN_DIR = ".chaoxing_grade_plans"
DEFAULT_REVIEW_BUNDLE_DIR = ".chaoxing_review_bundles"
DEFAULT_QUESTION_DRAFT_DIR = ".chaoxing_question_drafts"
DEFAULT_LOCAL_DEPS_DIR = ".chaoxing_deps"
QUESTION_TYPE_ALIASES = {
    "single_choice": "single_choice",
    "single": "single_choice",
    "choice": "single_choice",
    "单选": "single_choice",
    "单选题": "single_choice",
    "multiple_choice": "multiple_choice",
    "multiple": "multiple_choice",
    "multi": "multiple_choice",
    "多选": "multiple_choice",
    "多选题": "multiple_choice",
    "true_false": "true_false",
    "truefalse": "true_false",
    "tf": "true_false",
    "判断": "true_false",
    "判断题": "true_false",
    "short_answer": "short_answer",
    "short": "short_answer",
    "问答": "short_answer",
    "问答题": "short_answer",
    "简答": "short_answer",
    "简答题": "short_answer",
    "essay": "essay",
    "论述": "essay",
    "论述题": "essay",
}
QUESTION_TYPE_NAMES = {
    "single_choice": "单选题",
    "multiple_choice": "多选题",
    "true_false": "判断题",
    "short_answer": "简答题",
    "essay": "论述题",
}
DEFAULT_QUESTION_SCORES = {
    "single_choice": 2.0,
    "multiple_choice": 3.0,
    "true_false": 1.0,
    "short_answer": 8.0,
    "essay": 15.0,
}
# Online exam creation/publishing is hidden until those flows are verified.
TASK_ORDER = ("homework", "exam")
TASK_DEFINITIONS = {
    "exam": {
        "index": 1,
        "name": "考试界面",
        "module": "ks",
        "page_header": 7,
        "data_url": "/mooc2-ans/exam/test",
    },
    "homework": {
        "index": 2,
        "name": "作业界面",
        "module": "zy",
        "page_header": 6,
        "data_url": "/mooc2-ans/work/list",
    },
}
TASK_ALIASES = {
    "exam": "exam",
    "test": "exam",
    "ks": "exam",
    "考试": "exam",
    "考试界面": "exam",
    "homework": "homework",
    "work": "homework",
    "assignment": "homework",
    "zy": "homework",
    "zuoye": "homework",
    "作业": "homework",
    "作业界面": "homework",
}

LOGIN_URL_HINTS = (
    "passport2.chaoxing.com",
    "passport.chaoxing.com",
    "login.chaoxing.com",
    "/login",
)
LOGIN_TEXT_HINTS = (
    "账号登录",
    "手机号登录",
    "密码登录",
    "扫码登录",
    "验证码",
    "登录",
    "login",
)
LOGGED_IN_TEXT_HINTS = (
    "退出",
    "个人空间",
    "学习空间",
    "我的课程",
    "进入空间",
)
LOGGED_OUT_TEXT_HINTS = (
    "忘记密码",
    "新用户注册",
    "验证码登录",
    "手机号/超星号",
)


class ToolError(RuntimeError):
    pass


def script_dir() -> str:
    return os.path.dirname(os.path.abspath(__file__))


def local_deps_python_tag() -> str:
    return f"{sys.implementation.name}{sys.version_info.major}.{sys.version_info.minor}"


def default_local_deps_dir() -> str:
    return os.path.join(script_dir(), DEFAULT_LOCAL_DEPS_DIR, local_deps_python_tag())


def default_requirements_file() -> str:
    return os.path.join(script_dir(), "requirements.txt")


def add_local_deps_to_sys_path() -> None:
    deps_dir = default_local_deps_dir()
    if os.path.isdir(deps_dir) and deps_dir not in sys.path:
        sys.path.insert(0, deps_dir)


add_local_deps_to_sys_path()


def run_browser(args: list[str], *, check: bool = True) -> subprocess.CompletedProcess[str]:
    command = ["agent-browser", *args]
    result = subprocess.run(
        command,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if check and result.returncode != 0:
        detail = result.stderr.strip() or result.stdout.strip()
        raise ToolError(f"agent-browser failed: {detail}")
    return result


def base_args(session_name: str, headed: bool) -> list[str]:
    args = ["--session", session_name]
    if headed:
        args.append("--headed")
    return args


def ensure_agent_browser() -> None:
    if not shutil.which("agent-browser"):
        raise ToolError(
            "agent-browser is not installed. Install it with: npm i -g agent-browser && agent-browser install"
        )


def default_cookie_file() -> str:
    return os.path.join(script_dir(), ".chaoxing_cookies.json")


def ensure_api_dependencies() -> None:
    try:
        import requests  # noqa: F401
        from cryptography.hazmat.primitives.ciphers import Cipher  # noqa: F401
    except ImportError as exc:
        raise ToolError(
            "API commands need requests and cryptography. Run: "
            "python3 chaoxing_portal_tool.py setup"
        ) from exc


def safe_find_spec(import_name: str):
    try:
        return importlib.util.find_spec(import_name), ""
    except (ImportError, AttributeError, ValueError) as exc:
        return None, str(exc)


def api_dependency_report() -> list[dict[str, object]]:
    modules = [
        ("requests", "requests", "HTTP API requests", True),
        (
            "cryptography",
            "cryptography.hazmat.primitives.ciphers",
            "Chaoxing API password encryption",
            True,
        ),
        ("pdfplumber", "pdfplumber", "PDF attachment text extraction", False),
    ]
    report = []
    for module_name, import_name, purpose, required in modules:
        spec, import_error = safe_find_spec(import_name)
        if spec is not None:
            try:
                __import__(import_name)
            except Exception as exc:
                import_error = str(exc)
        report.append(
            {
                "module": module_name,
                "available": spec is not None and not import_error,
                "purpose": purpose,
                "required": required,
                "origin": getattr(spec, "origin", "") if spec else "",
                "error": import_error,
            }
        )
    return report


def command_setup(args: argparse.Namespace) -> int:
    target_dir = os.path.abspath(args.target)
    requirements = os.path.abspath(args.requirements)
    if not os.path.exists(requirements):
        raise ToolError(f"requirements file not found: {requirements}")

    os.makedirs(target_dir, exist_ok=True)
    command = [
        sys.executable,
        "-m",
        "pip",
        "install",
        "-r",
        requirements,
        "--target",
        target_dir,
    ]
    if args.upgrade:
        command.append("--upgrade")

    print(f"python: {sys.executable}")
    print(f"python_version: {sys.version.split()[0]}")
    print(f"requirements: {requirements}")
    print(f"target: {target_dir}")
    result = subprocess.run(command, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if result.returncode != 0:
        if result.stdout.strip():
            print(result.stdout.strip())
        if result.stderr.strip():
            print("pip_error:")
            print(result.stderr.strip())
        raise ToolError(
            "dependency install failed for this Python. Use the same working Python for doctor, setup, "
            "login, and grading commands. If pip itself is broken, repair/reinstall Python or switch to "
            "another Python executable, then run that Python with chaoxing_portal_tool.py setup."
        )

    add_local_deps_to_sys_path()
    if target_dir not in sys.path:
        sys.path.insert(0, target_dir)
    importlib.invalidate_caches()
    missing = [row["module"] for row in api_dependency_report() if not row["available"] and row["required"]]
    if missing:
        raise ToolError("installed dependencies, but still missing: " + ", ".join(missing))

    print("setup: ok")
    print("You can now run: python3 chaoxing_portal_tool.py doctor")
    return 0


def command_doctor(args: argparse.Namespace) -> int:
    print(f"tool_dir: {script_dir()}")
    print(f"python: {sys.executable}")
    print(f"python_version: {sys.version.split()[0]}")
    print(f"local_deps: {default_local_deps_dir()}")
    print(f"requirements: {default_requirements_file()}")

    problems = []
    for row in api_dependency_report():
        status = "ok" if row["available"] else "missing"
        required = "required" if row["required"] else "recommended"
        print(f"dependency:{row['module']}: {status} ({required})")
        if row["origin"]:
            print(f"  origin: {row['origin']}")
        if row.get("error"):
            print(f"  error: {row['error']}")
        if not row["available"] and row["required"]:
            problems.append(f"missing Python package: {row['module']}")
        elif not row["available"]:
            print("  note: optional; run setup if you want richer attachment extraction")

    cookie_file = args.cookie_file
    cookie_status = "present" if os.path.exists(cookie_file) else "missing"
    print(f"cookie_file: {cookie_file}")
    print(f"cookie_status: {cookie_status}")

    agent_browser = shutil.which("agent-browser")
    print(f"agent-browser: {agent_browser or 'missing'}")
    if not agent_browser:
        print("  browser commands are optional; API homework workflows do not need agent-browser")

    if args.check_login:
        try:
            ok, message = saved_login_status(cookie_file, args.check_url, args.timeout)
        except Exception as exc:
            ok, message = False, str(exc)
        print(f"login_check: {'ok' if ok else 'failed'}")
        print(f"login_message: {message}")
        if not ok:
            problems.append("saved Chaoxing login is unavailable")

    if problems:
        print("doctor: needs_attention")
        if any(problem.startswith("missing Python package") for problem in problems):
            print("next_step: python3 chaoxing_portal_tool.py setup")
        elif any("login" in problem for problem in problems):
            print("next_step: python3 chaoxing_portal_tool.py login")
        return 2

    print("doctor: ok")
    return 0


def request_headers() -> dict[str, str]:
    return {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    }


def aes_encrypt_base64(value: str, key: str = TRANSFER_KEY) -> str:
    from cryptography.hazmat.primitives import padding
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes

    key_bytes = key.encode("utf-8")
    padder = padding.PKCS7(algorithms.AES.block_size).padder()
    padded = padder.update(value.encode("utf-8")) + padder.finalize()
    encryptor = Cipher(algorithms.AES(key_bytes), modes.CBC(key_bytes)).encryptor()
    encrypted = encryptor.update(padded) + encryptor.finalize()
    return base64.b64encode(encrypted).decode("ascii")


def parse_hidden_inputs(html: str) -> dict[str, str]:
    hidden: dict[str, str] = {}
    for match in re.finditer(r"<input\b[^>]*>", html, flags=re.I | re.S):
        tag = match.group(0)
        attrs = html_attrs(tag)
        if attrs.get("type", "").lower() != "hidden":
            continue
        key = attrs.get("name") or attrs.get("id")
        if key:
            hidden[key] = attrs.get("value", "")
    return hidden


def html_attrs(tag: str) -> dict[str, str]:
    return {
        name.lower(): unescape(value)
        for name, _quote, value in re.findall(
            r"([a-zA-Z_:][\w:.-]*)\s*=\s*(['\"])(.*?)\2",
            tag,
            flags=re.S,
        )
    }


def cookie_file_payload(session: object) -> dict[str, object]:
    cookies = []
    for cookie in session.cookies:
        cookies.append(
            {
                "name": cookie.name,
                "value": cookie.value,
                "domain": cookie.domain,
                "path": cookie.path,
                "secure": bool(cookie.secure),
                "expires": cookie.expires,
                "httpOnly": cookie.has_nonstandard_attr("HttpOnly"),
            }
        )
    return {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "cookies": cookies,
    }


def save_cookie_file(session: object, path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    if directory:
        os.makedirs(directory, exist_ok=True)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(cookie_file_payload(session), file, ensure_ascii=False, indent=2)
        file.write("\n")
    os.chmod(path, 0o600)


def load_cookie_file(path: str) -> list[dict[str, object]]:
    if not os.path.exists(path):
        raise ToolError(f"cookie file not found: {path}")
    with open(path, "r", encoding="utf-8") as file:
        payload = json.load(file)
    cookies = payload.get("cookies", [])
    if not isinstance(cookies, list):
        raise ToolError(f"invalid cookie file: {path}")
    return cookies


def request_session_from_cookie_file(path: str):
    import requests

    session = requests.Session()
    session.headers.update(request_headers())
    for cookie in load_cookie_file(path):
        name = str(cookie.get("name", ""))
        value = str(cookie.get("value", ""))
        domain = str(cookie.get("domain", ""))
        path_value = str(cookie.get("path", "/"))
        if name and domain:
            session.cookies.set(name, value, domain=domain, path=path_value)
    return session


def inject_cookie_file_into_browser(browser_args: list[str], path: str) -> int:
    if not os.path.exists(path):
        return 0

    injected = 0
    for cookie in load_cookie_file(path):
        name = str(cookie.get("name", ""))
        value = str(cookie.get("value", ""))
        domain = str(cookie.get("domain", ""))
        path_value = str(cookie.get("path", "/"))
        if not name or not domain:
            continue

        cookie_url = "https://" + domain.lstrip(".") + (path_value if path_value.startswith("/") else "/" + path_value)
        command = [
            *browser_args,
            "cookies",
            "set",
            name,
            value,
            "--url",
            cookie_url,
            "--domain",
            domain,
            "--path",
            path_value,
        ]
        if cookie.get("secure"):
            command.append("--secure")
        if cookie.get("httpOnly"):
            command.append("--httpOnly")
        expires = cookie.get("expires")
        if expires:
            command.extend(["--expires", str(expires)])
        if run_browser(command, check=False).returncode == 0:
            injected += 1
    return injected


def current_page_summary(browser_args: list[str]) -> tuple[str, str, str]:
    url = run_browser([*browser_args, "get", "url"]).stdout.strip()
    title = run_browser([*browser_args, "get", "title"], check=False).stdout.strip()
    body = run_browser([*browser_args, "get", "text", "body"], check=False).stdout
    return url, title, body


def page_summary(session_name: str, headed: bool) -> tuple[str, str, str]:
    return current_page_summary(base_args(session_name, headed))


def looks_logged_out(url: str, title: str, body: str) -> bool:
    haystack_url = url.lower()
    if any(hint in haystack_url for hint in LOGIN_URL_HINTS):
        return True

    sample = f"{title}\n{body[:12000]}".lower()
    if any(hint.lower() in sample for hint in LOGGED_IN_TEXT_HINTS):
        return False
    if any(hint.lower() in sample for hint in LOGGED_OUT_TEXT_HINTS):
        return True
    return any(hint.lower() in sample for hint in LOGIN_TEXT_HINTS)


def print_page_status(session_name: str, url: str, title: str, body: str) -> int:
    logged_out = looks_logged_out(url, title, body)

    print(f"session: {session_name}")
    print(f"url: {url}")
    print(f"title: {title or '(no title found)'}")
    print(f"auth_result: {'logged_out_or_login_page' if logged_out else 'probably_logged_in'}")
    return 3 if logged_out else 0


def auth_code_from_current_page(browser_args: list[str]) -> int:
    url, title, body = current_page_summary(browser_args)
    return 3 if looks_logged_out(url, title, body) else 0


def print_status(session_name: str, headed: bool) -> int:
    url, title, body = page_summary(session_name, headed)
    return print_page_status(session_name, url, title, body)


def prompt_credentials() -> tuple[str, str]:
    username = input("Chaoxing username / phone: ").strip()
    if not username:
        raise ToolError("username cannot be empty")

    password = getpass.getpass("Chaoxing password: ")
    if not password:
        raise ToolError("password cannot be empty")

    return username, password


def apple_script_quote(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def macos_dialog_text(prompt: str, title: str, hidden: bool = False) -> str:
    command = [
        "osascript",
        "-e",
        (
            "text returned of (display dialog "
            + apple_script_quote(prompt)
            + " default answer \"\" "
            + ("with hidden answer " if hidden else "")
            + "with title "
            + apple_script_quote(title)
            + " buttons {\"Cancel\", \"OK\"} default button \"OK\")"
        ),
    ]
    result = subprocess.run(command, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if result.returncode != 0:
        raise ToolError("login prompt was canceled or unavailable")
    return result.stdout.strip()


def prompt_credentials_macos_dialog() -> tuple[str, str]:
    username = macos_dialog_text("请输入超星账号/手机号", "Chaoxing Login")
    if not username:
        raise ToolError("username cannot be empty")

    password = macos_dialog_text("请输入超星密码", "Chaoxing Login", hidden=True)
    if not password:
        raise ToolError("password cannot be empty")

    return username, password


def prompt_credentials_stdin_json() -> tuple[str, str]:
    try:
        payload = json.load(sys.stdin)
    except Exception as exc:
        raise ToolError("--stdin-json expects a JSON object on stdin") from exc
    if not isinstance(payload, dict):
        raise ToolError("--stdin-json expects a JSON object on stdin")

    username = str(payload.get("username") or payload.get("phone") or payload.get("account") or "").strip()
    password = str(payload.get("password") or "")
    if not username:
        raise ToolError("stdin JSON is missing username/phone/account")
    if not password:
        raise ToolError("stdin JSON is missing password")
    return username, password


def open_login_dialog(browser_args: list[str]) -> None:
    snapshot = run_browser(
        [*browser_args, "snapshot", "-i", "-u", "-d", "4"],
        check=False,
    ).stdout
    if "手机号/超星号" in snapshot or "密码" in snapshot:
        return

    if "登录" in snapshot:
        clicked = run_browser(
            [*browser_args, "find", "role", "button", "click", "--name", "登录"],
            check=False,
        )
        if clicked.returncode != 0:
            run_browser([*browser_args, "find", "text", "登录", "click", "--exact"], check=False)
        run_browser([*browser_args, "wait", "--text", "手机号/超星号"], check=False)


def fill_and_submit_password_login(
    browser_args: list[str],
    url: str,
    username: str,
    password: str,
) -> int:
    run_browser([*browser_args, "open", url])
    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
    open_login_dialog(browser_args)

    run_browser([*browser_args, "find", "text", "手机号登录", "click"], check=False)
    username_filled = run_browser(
        [*browser_args, "find", "placeholder", "手机号/超星号", "fill", username],
        check=False,
    )
    if username_filled.returncode != 0:
        run_browser([*browser_args, "fill", "input[type=text]", username], check=False)

    password_filled = run_browser(
        [*browser_args, "find", "placeholder", "密码", "fill", password],
        check=False,
    )
    if password_filled.returncode != 0:
        run_browser([*browser_args, "fill", "input[type=password]", password], check=False)

    submitted = run_browser(
        [*browser_args, "find", "role", "button", "click", "--name", "登录"],
        check=False,
    )
    if submitted.returncode != 0:
        run_browser([*browser_args, "press", "Enter"], check=False)

    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
    code = print_status_from_current_page(browser_args)
    if code != 0:
        print("If the browser asks for captcha, SMS, or QR confirmation, complete it there.")
        print("After the portal page finishes loading, press Enter here to save/check the session.")
        input()
        run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
        return print_status_from_current_page(browser_args)
    return code


def manual_browser_login(
    browser_args: list[str],
    url: str,
    timeout: float,
    interval: float,
) -> int:
    interval = max(interval, 0.25)
    run_browser([*browser_args, "open", url])
    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)

    session_name = browser_args[browser_args.index("--session-name") + 1]
    print(f"Cookies and browser storage will be saved under session: {session_name}")

    code = print_status_from_current_page(browser_args)
    if code == 0:
        print("The saved session is already logged in.")
        return 0

    open_login_dialog(browser_args)
    print("A browser login window has opened.")
    print("Complete the normal Chaoxing/Fanya login flow in that window.")

    if timeout > 0:
        print(f"Waiting up to {timeout:g}s for login to finish...")
        deadline = time.time() + timeout
        while time.time() < deadline:
            time.sleep(interval)
            run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
            if auth_code_from_current_page(browser_args) == 0:
                return print_status_from_current_page(browser_args)

    print("After the portal page finishes loading, return here and press Enter.")
    input()
    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
    return print_status_from_current_page(browser_args)


def print_status_from_current_page(browser_args: list[str]) -> int:
    url, title, body = current_page_summary(browser_args)
    session_name = browser_args[browser_args.index("--session-name") + 1]
    return print_page_status(session_name, url, title, body)


def api_login(args: argparse.Namespace) -> int:
    ensure_api_dependencies()
    import requests

    username = os.environ.get("CHAOXING_USERNAME", "").strip()
    password = os.environ.get("CHAOXING_PASSWORD", "")
    if getattr(args, "stdin_json", False):
        username, password = prompt_credentials_stdin_json()
    elif getattr(args, "macos_dialog", False):
        username, password = prompt_credentials_macos_dialog()
    elif not username or args.prompt:
        username, password = prompt_credentials()
    elif not password:
        password = getpass.getpass("Chaoxing password: ")
        if not password:
            raise ToolError("password cannot be empty")

    session = requests.Session()
    session.headers.update(request_headers())

    refer = quote(args.url, safe="")
    login_page = f"{LOGIN_PAGE_URL}?newversion=true&fid={args.fid}&refer={refer}"
    try:
        page_response = session.get(login_page, timeout=args.timeout)
        page_response.raise_for_status()
    except requests.RequestException as exc:
        raise ToolError(f"login page request failed: {exc}") from exc
    hidden = parse_hidden_inputs(page_response.text)

    t_value = hidden.get("t", "true")
    post_username = username
    post_password = password
    if t_value == "true":
        post_username = aes_encrypt_base64(username)
        post_password = aes_encrypt_base64(password)

    payload = {
        "fid": hidden.get("fid", args.fid),
        "uname": post_username,
        "password": post_password,
        "refer": hidden.get("refer", refer),
        "t": t_value,
        "forbidotherlogin": hidden.get("forbidotherlogin", "0"),
        "validate": hidden.get("validate", ""),
        "doubleFactorLogin": hidden.get("doubleFactorLogin", "0"),
        "independentId": hidden.get("independentId", ""),
        "independentNameId": hidden.get("independentNameId", ""),
    }

    try:
        response = session.post(
            FANYA_LOGIN_URL,
            data=payload,
            timeout=args.timeout,
            headers={
                **request_headers(),
                "Origin": "https://passport2.chaoxing.com",
                "Referer": login_page,
                "X-Requested-With": "XMLHttpRequest",
                "Accept": "application/json, text/javascript, */*; q=0.01",
            },
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise ToolError(f"login request failed: {exc}") from exc
    try:
        result = response.json()
    except ValueError as exc:
        raise ToolError(f"login response was not JSON: {response.text[:200]}") from exc

    if result.get("containTwoFactorLogin"):
        two_factor_url = result.get("twoFactorLoginPCUrl", "")
        raise ToolError(f"login requires two-factor verification: {two_factor_url}")

    if not result.get("status"):
        message = result.get("msg2") or result.get("mes") or result.get("errorMsg") or "login failed"
        raise ToolError(str(message))

    target_url = unquote(str(result.get("url") or args.url))
    try:
        if target_url:
            session.get(target_url, timeout=args.timeout, allow_redirects=True)
        session.get(args.url, timeout=args.timeout, allow_redirects=True)
    except requests.RequestException as exc:
        raise ToolError(f"post-login cookie warmup failed: {exc}") from exc

    save_cookie_file(session, args.cookie_file)
    print("api_login: success")
    print(f"cookie_file: {args.cookie_file}")
    print(f"cookies_saved: {len(session.cookies)}")
    if not getattr(args, "check_after", True):
        return 0
    return api_status(
        argparse.Namespace(
            cookie_file=args.cookie_file,
            check_url=args.check_url,
            timeout=args.timeout,
            auto_login=False,
            url=args.url,
        )
    )


def saved_login_status(cookie_file: str, check_url: str, timeout: float) -> tuple[bool, str]:
    ensure_api_dependencies()
    try:
        session = request_session_from_cookie_file(cookie_file)
    except ToolError as exc:
        return False, str(exc)

    try:
        response = session.get(check_url, timeout=timeout, allow_redirects=True)
    except Exception as exc:
        return False, f"status request failed: {exc}"

    title_match = re.search(r"<title[^>]*>(.*?)</title>", response.text, flags=re.I | re.S)
    title = re.sub(r"\s+", " ", title_match.group(1)).strip() if title_match else ""
    logged_out = looks_logged_out(response.url, title, response.text)
    if logged_out:
        return False, f"saved cookies are logged out: {response.url}"
    return True, f"probably logged in: {response.url}"


def ensure_api_logged_in(args: argparse.Namespace) -> None:
    check_url = getattr(args, "check_url", AUTH_CHECK_URL)
    timeout = getattr(args, "login_timeout", getattr(args, "timeout", 20.0))
    ok, message = saved_login_status(args.cookie_file, check_url, timeout)
    if ok:
        return

    print(f"saved login unavailable: {message}")
    print("Please log in with Chaoxing account/password.")
    api_login(
        argparse.Namespace(
            cookie_file=args.cookie_file,
            check_url=check_url,
            timeout=timeout,
            url=args.url,
            fid=getattr(args, "fid", "-1"),
            prompt=True,
            check_after=False,
        )
    )


def api_status(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)

    ensure_api_dependencies()
    session = request_session_from_cookie_file(args.cookie_file)
    try:
        response = session.get(args.check_url, timeout=args.timeout, allow_redirects=True)
    except Exception as exc:
        raise ToolError(f"status request failed: {exc}") from exc

    title_match = re.search(r"<title[^>]*>(.*?)</title>", response.text, flags=re.I | re.S)
    title = re.sub(r"\s+", " ", title_match.group(1)).strip() if title_match else ""
    logged_out = looks_logged_out(response.url, title, response.text)

    print(f"cookie_file: {args.cookie_file}")
    print(f"status: {response.status_code}")
    print(f"final_url: {response.url}")
    print(f"title: {unescape(title) or '(no title found)'}")
    print(f"auth_result: {'logged_out_or_login_page' if logged_out else 'probably_logged_in'}")
    return 3 if logged_out else 0


def fetch_backclazz_data(cookie_file: str, timeout: float) -> dict[str, object]:
    ensure_api_dependencies()
    session = request_session_from_cookie_file(cookie_file)
    try:
        response = session.get(BACKCLAZZ_URL, timeout=timeout)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"course list request failed: {exc}") from exc
    try:
        data = response.json()
    except ValueError as exc:
        raise ToolError(f"course list response was not JSON: {response.text[:200]}") from exc
    if data.get("result") != 1:
        raise ToolError(str(data.get("msg") or "failed to fetch course list"))
    return data


def detect_current_user_name(cookie_file: str, timeout: float) -> str:
    ensure_api_dependencies()
    session = request_session_from_cookie_file(cookie_file)
    try:
        response = session.get(AUTH_CHECK_URL, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
    except Exception:
        return ""

    html = response.text
    patterns = (
        r'aria-label="账号：([^"]+)"',
        r'<li class="user">\s*<img[^>]*>\s*<h3[^>]*>([^<]+)</h3>',
    )
    for pattern in patterns:
        match = re.search(pattern, html, flags=re.I | re.S)
        if match:
            return unescape(match.group(1)).strip()
    return ""


def teaching_courses_from_backclazz(data: dict[str, object]) -> list[dict[str, object]]:
    courses: list[dict[str, object]] = []
    seen: set[str] = set()
    for channel in data.get("channelList", []):
        if not isinstance(channel, dict):
            continue
        content = channel.get("content") or {}
        if not isinstance(content, dict):
            continue
        clazzes = content.get("clazz")
        if not isinstance(clazzes, list) or not clazzes:
            continue

        course_id = str(content.get("id") or "").strip()
        if not course_id or course_id in seen:
            continue
        seen.add(course_id)

        courses.append(
            {
                "index": len(courses) + 1,
                "course_id": course_id,
                "course_name": str(content.get("name") or "").strip(),
                "teacher": str(content.get("teacherfactor") or "").strip(),
                "term": str(content.get("schools") or "").strip(),
                "cpi": str(channel.get("cpi") or content.get("cpi") or "").strip(),
                "class_count": len(clazzes),
                "classes": [
                    {
                        "index": idx + 1,
                        "clazz_id": str(clazz.get("clazzId") or "").strip(),
                        "clazz_name": str(clazz.get("clazzName") or "").strip(),
                        "student_count": clazz.get("clazzStudentCount"),
                        "chatid": str(clazz.get("chatid") or "").strip(),
                        "state": clazz.get("state"),
                    }
                    for idx, clazz in enumerate(clazzes)
                    if isinstance(clazz, dict)
                ],
            }
        )
    return courses


def prioritize_user_classes(course: dict[str, object], owner_name: str) -> None:
    if not owner_name:
        return

    classes = course.get("classes", [])
    if not isinstance(classes, list):
        return

    for clazz in classes:
        if isinstance(clazz, dict):
            clazz["is_user_class"] = owner_name in str(clazz.get("clazz_name", ""))

    classes.sort(key=lambda clazz: 0 if clazz.get("is_user_class") else 1)
    for index, clazz in enumerate(classes, 1):
        clazz["index"] = index
    course["current_user"] = owner_name


def find_teaching_course(courses: list[dict[str, object]], query: str) -> dict[str, object]:
    normalized = query.strip().lower()
    if not normalized:
        raise ToolError("course name/id cannot be empty")

    exact = [
        course
        for course in courses
        if str(course.get("course_id", "")).lower() == normalized
        or str(course.get("course_name", "")).strip().lower() == normalized
    ]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        names = ", ".join(str(course.get("course_name")) for course in exact)
        raise ToolError(f"multiple exact course matches: {names}")

    partial = [
        course
        for course in courses
        if normalized in str(course.get("course_name", "")).strip().lower()
    ]
    if len(partial) == 1:
        return partial[0]
    if not partial:
        raise ToolError(f"course not found: {query}")

    names = ", ".join(f"{course['index']}. {course['course_name']}" for course in partial)
    raise ToolError(f"multiple course matches: {names}")


def find_course_class(course: dict[str, object], query: str) -> dict[str, object]:
    normalized = query.strip().lower()
    if not normalized:
        raise ToolError("class name/index/clazz_id cannot be empty")

    classes = course.get("classes", [])
    if not isinstance(classes, list):
        raise ToolError("selected course has no class list")

    exact_id = [
        clazz
        for clazz in classes
        if isinstance(clazz, dict) and str(clazz.get("clazz_id", "")).strip().lower() == normalized
    ]
    if len(exact_id) == 1:
        return exact_id[0]

    exact_index = [
        clazz
        for clazz in classes
        if isinstance(clazz, dict) and str(clazz.get("index", "")).strip().lower() == normalized
    ]
    if len(exact_index) == 1:
        return exact_index[0]

    exact_name = [
        clazz
        for clazz in classes
        if isinstance(clazz, dict) and str(clazz.get("clazz_name", "")).strip().lower() == normalized
    ]
    if len(exact_name) == 1:
        return exact_name[0]
    if len(exact_name) > 1:
        names = ", ".join(str(clazz.get("clazz_name")) for clazz in exact_name)
        raise ToolError(f"multiple exact class matches: {names}")

    partial = [
        clazz
        for clazz in classes
        if isinstance(clazz, dict) and normalized in str(clazz.get("clazz_name", "")).strip().lower()
    ]
    if len(partial) == 1:
        return partial[0]
    if not partial:
        raise ToolError(f"class not found in {course.get('course_name')}: {query}")

    names = ", ".join(f"{clazz['index']}. {clazz['clazz_name']}" for clazz in partial)
    raise ToolError(f"multiple class matches: {names}")


def normalize_task_kind(task: str) -> str:
    normalized = task.strip().lower()
    if not normalized:
        raise ToolError("task cannot be empty")

    task_kind = TASK_ALIASES.get(normalized)
    if not task_kind:
        raise ToolError("task must be one of: homework/作业, exam/考试")
    return task_kind


def task_options() -> list[dict[str, object]]:
    return [
        {
            "index": index,
            "task": task_kind,
            "name": TASK_DEFINITIONS[task_kind]["name"],
            "module": TASK_DEFINITIONS[task_kind]["module"],
            "page_header": TASK_DEFINITIONS[task_kind]["page_header"],
        }
        for index, task_kind in enumerate(TASK_ORDER, 1)
    ]


def course_summary(course: dict[str, object]) -> dict[str, object]:
    return {
        "course_id": course.get("course_id", ""),
        "course_name": course.get("course_name", ""),
        "teacher": course.get("teacher", ""),
        "term": course.get("term", ""),
        "cpi": course.get("cpi", ""),
    }


def class_summary(clazz: dict[str, object]) -> dict[str, object]:
    return {
        "index": clazz.get("index", ""),
        "clazz_id": clazz.get("clazz_id", ""),
        "clazz_name": clazz.get("clazz_name", ""),
        "student_count": clazz.get("student_count"),
        "is_user_class": bool(clazz.get("is_user_class")),
    }


def resolve_course_and_class(args: argparse.Namespace) -> tuple[dict[str, object], dict[str, object]]:
    courses = teaching_courses_from_backclazz(fetch_backclazz_data(args.cookie_file, args.timeout))
    course = find_teaching_course(courses, args.course)
    owner_name = args.owner_name or detect_current_user_name(args.cookie_file, args.timeout)
    prioritize_user_classes(course, owner_name)
    clazz = find_course_class(course, args.clazz)
    return course, clazz


def build_teacher_entry_url(course: dict[str, object], clazz: dict[str, object], page_header: int) -> str:
    course_id = str(course.get("course_id") or "").strip()
    clazz_id = str(clazz.get("clazz_id") or "").strip()
    cpi = str(course.get("cpi") or "").strip()
    if not course_id or not clazz_id or not cpi:
        raise ToolError("course_id, clazz_id, and cpi are required to enter course tasks")

    return COURSE_ENTRY_URL + "?" + urlencode(
        {
            "courseId": course_id,
            "clazzId": clazz_id,
            "edit": "true",
            "v": "2",
            "cpi": cpi,
            "pageHeader": str(page_header),
            "single": "0",
        }
    )


def parse_teacher_nav_items(html: str) -> dict[str, dict[str, str]]:
    items: dict[str, dict[str, str]] = {}
    for match in re.finditer(r"<li\b([^>]*)>(.*?)</li>", html, flags=re.I | re.S):
        li_attrs = html_attrs(match.group(1))
        module = li_attrs.get("dataname", "")
        if not module:
            continue

        body = match.group(2)
        anchor = re.search(r"<a\b[^>]*>", body, flags=re.I | re.S)
        anchor_attrs = html_attrs(anchor.group(0)) if anchor else {}
        data_url = anchor_attrs.get("data-url", "")
        if not data_url:
            continue

        items[module] = {
            "module": module,
            "title": anchor_attrs.get("title", ""),
            "data_url": data_url,
            "page_header": li_attrs.get("pageheader", ""),
            "course_nav_id": li_attrs.get("data", ""),
            "open_type": li_attrs.get("opentype", ""),
        }
    return items


def fetch_teacher_task_context(
    cookie_file: str,
    course: dict[str, object],
    clazz: dict[str, object],
    page_header: int,
    timeout: float,
) -> dict[str, object]:
    ensure_api_dependencies()
    session = request_session_from_cookie_file(cookie_file)
    entry_url = build_teacher_entry_url(course, clazz, page_header)
    try:
        response = session.get(entry_url, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"teacher course page request failed: {exc}") from exc

    title_match = re.search(r"<title[^>]*>(.*?)</title>", response.text, flags=re.I | re.S)
    title = re.sub(r"\s+", " ", title_match.group(1)).strip() if title_match else ""
    if looks_logged_out(response.url, title, response.text):
        raise ToolError(f"saved cookies are logged out: {response.url}")

    hidden = parse_hidden_inputs(response.text)
    nav_items = parse_teacher_nav_items(response.text)
    for required in ("courseid", "clazzid", "cpi", "enc", "openc", "t"):
        if not hidden.get(required):
            raise ToolError(f"teacher course page is missing required field: {required}")

    return {
        "entry_url": entry_url,
        "final_url": response.url,
        "hidden": hidden,
        "nav_items": nav_items,
    }


def build_task_content_url(context: dict[str, object], task_kind: str, clazz: dict[str, object]) -> str:
    task = TASK_DEFINITIONS[task_kind]
    hidden = context.get("hidden")
    nav_items = context.get("nav_items")
    if not isinstance(hidden, dict) or not isinstance(nav_items, dict):
        raise ToolError("invalid teacher task context")

    module = str(task["module"])
    nav_item = nav_items.get(module, {})
    if not isinstance(nav_item, dict):
        nav_item = {}
    data_url = str(nav_item.get("data_url") or task["data_url"])
    base_url = urljoin(str(context.get("final_url", "")), data_url)
    separator = "&" if "?" in base_url else "?"
    clazz_id = str(clazz.get("clazz_id") or hidden.get("clazzid") or "").strip()

    params = {
        "courseid": str(hidden.get("courseid", "")),
        "clazzid": clazz_id,
        "courseId": str(hidden.get("courseid", "")),
        "classId": clazz_id,
        "clazzId": clazz_id,
        "cpi": str(hidden.get("cpi", "")),
        "enc": str(hidden.get("enc", "")),
        "openc": str(hidden.get("openc", "")),
        "t": str(hidden.get("t", "")),
        "ut": "t",
    }
    if task_kind == "homework":
        params["selectClassid"] = clazz_id
    return base_url + separator + urlencode(params)


def append_query_params(url: str, params: dict[str, object]) -> str:
    return url + ("&" if "?" in url else "?") + urlencode(params)


def compact_text(fragment: str) -> str:
    fragment = re.sub(r"<script\b.*?</script>", " ", fragment, flags=re.I | re.S)
    fragment = re.sub(r"<style\b.*?</style>", " ", fragment, flags=re.I | re.S)
    fragment = re.sub(r"<br\s*/?>", "\n", fragment, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", fragment)
    text = unescape(text)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text)
    return text.strip()


def query_value(url: str, name: str) -> str:
    values = parse_qs(urlparse(url).query).get(name, [])
    return values[0] if values else ""


def fetch_homework_list_page(
    cookie_file: str,
    course: dict[str, object],
    clazz: dict[str, object],
    page: int,
    timeout: float,
) -> tuple[str, str]:
    task = TASK_DEFINITIONS["homework"]
    context = fetch_teacher_task_context(cookie_file, course, clazz, int(task["page_header"]), timeout)
    list_url = build_task_content_url(context, "homework", clazz)
    list_url = append_query_params(
        list_url,
        {
            "pages": str(page),
            "pageSize": "12",
            "status": "-1",
        },
    )

    session = request_session_from_cookie_file(cookie_file)
    try:
        response = session.get(list_url, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"homework list request failed: {exc}") from exc
    return response.text, response.url


def parse_homework_items(html: str, page_url: str, clazz: dict[str, object]) -> list[dict[str, object]]:
    items: list[dict[str, object]] = []
    class_name = str(clazz.get("clazz_name") or "")
    for match in re.finditer(r"<li\b[^>]*id=['\"]work\d+['\"][\s\S]*?</li>", html, flags=re.I):
        block = match.group(0)
        open_tag = re.search(r"<li\b[^>]*>", block, flags=re.I)
        attrs = html_attrs(open_tag.group(0)) if open_tag else {}
        plain = compact_text(block)

        mark_url_match = re.search(r"href=['\"]([^'\"]*?/mooc2-ans/work/mark\?[^'\"]+)['\"]", block, flags=re.I)
        mark_url = urljoin(page_url, unescape(mark_url_match.group(1))) if mark_url_match else ""
        work_id = query_value(mark_url, "id") or str(attrs.get("id", "")).removeprefix("work")
        if not work_id:
            continue

        counts = re.search(r"(\d+)\s*待批\s*(\d+)\s*已交\s*(\d+)\s*未交", plain)
        pending = int(counts.group(1)) if counts else 0
        submitted = int(counts.group(2)) if counts else None
        unsubmitted = int(counts.group(3)) if counts else None

        title_source = plain
        if class_name and class_name in title_source:
            title = title_source.split(class_name, 1)[0].strip()
        elif "作答时间" in title_source:
            title = title_source.split("作答时间", 1)[0].strip()
        else:
            title = re.sub(r"\d+\s*待批.*$", "", title_source).strip()

        time_match = re.search(r"作答时间[:：]\s*(.*?)\s*(?:修改设置|移动到|删除|批阅|$)", plain)
        answer_time = time_match.group(1).strip() if time_match else ""

        items.append(
            {
                "work_id": work_id,
                "title": title,
                "pending_count": pending,
                "submitted_count": submitted,
                "unsubmitted_count": unsubmitted,
                "answer_time": answer_time,
                "mark_url": mark_url,
            }
        )
    return items


def fetch_homework_items_for_class(
    cookie_file: str,
    course: dict[str, object],
    clazz: dict[str, object],
    timeout: float,
) -> list[dict[str, object]]:
    first_html, first_url = fetch_homework_list_page(cookie_file, course, clazz, 1, timeout)
    hidden = parse_hidden_inputs(first_html)
    total_pages = int(hidden.get("pageNum") or "1")
    items = parse_homework_items(first_html, first_url, clazz)
    for page in range(2, total_pages + 1):
        html, url = fetch_homework_list_page(cookie_file, course, clazz, page, timeout)
        items.extend(parse_homework_items(html, url, clazz))

    for index, item in enumerate(items, 1):
        item["index"] = index
    return items


def ungraded_homework_items(items: list[dict[str, object]]) -> list[dict[str, object]]:
    selected = [item for item in items if int(item.get("pending_count") or 0) > 0]
    for index, item in enumerate(selected, 1):
        item["index"] = index
    return selected


def find_homework_item(items: list[dict[str, object]], query: str) -> dict[str, object]:
    normalized = query.strip().lower()
    if not normalized:
        raise ToolError("homework index/title/work_id cannot be empty")

    exact_id = [
        item
        for item in items
        if str(item.get("work_id", "")).strip().lower() == normalized
    ]
    if len(exact_id) == 1:
        return exact_id[0]

    exact_index = [
        item
        for item in items
        if str(item.get("index", "")).strip().lower() == normalized
    ]
    if len(exact_index) == 1:
        return exact_index[0]

    exact_title = [
        item
        for item in items
        if str(item.get("title", "")).strip().lower() == normalized
    ]
    if len(exact_title) == 1:
        return exact_title[0]
    if len(exact_title) > 1:
        titles = ", ".join(str(item.get("title")) for item in exact_title)
        raise ToolError(f"multiple exact homework matches: {titles}")

    partial = [
        item
        for item in items
        if normalized in str(item.get("title", "")).strip().lower()
    ]
    if len(partial) == 1:
        return partial[0]
    if not partial:
        raise ToolError(f"homework not found: {query}")

    titles = ", ".join(f"{item['index']}. {item['title']}" for item in partial)
    raise ToolError(f"multiple homework matches: {titles}")


def fetch_work_info(
    session: object,
    course: dict[str, object],
    clazz: dict[str, object],
    work_id: str,
    timeout: float,
) -> dict[str, object]:
    try:
        response = session.get(
            MOOC2_BASE_URL + "/mooc2-ans/work/workinfo",
            params={
                "courseid": str(course.get("course_id") or ""),
                "clazzid": str(clazz.get("clazz_id") or ""),
                "workid": work_id,
                "cpi": str(course.get("cpi") or ""),
            },
            timeout=timeout,
        )
        response.raise_for_status()
        result = response.json()
    except Exception as exc:
        raise ToolError(f"work info request failed: {exc}") from exc

    if not result.get("status"):
        raise ToolError(str(result.get("msg") or "failed to fetch work info"))
    data = result.get("data")
    return data if isinstance(data, dict) else {}


def fetch_mark_list_page(
    session: object,
    course: dict[str, object],
    clazz: dict[str, object],
    work_id: str,
    status: int,
    page: int,
    size: int,
    timeout: float,
) -> str:
    try:
        response = session.get(
            MOOC2_BASE_URL + "/mooc2-ans/work/mark-list",
            params={
                "courseid": str(course.get("course_id") or ""),
                "clazzid": str(clazz.get("clazz_id") or ""),
                "workid": work_id,
                "submit": "true",
                "status": str(status),
                "groupId": "0",
                "cpi": str(course.get("cpi") or ""),
                "evaluation": "0",
                "sort": "0",
                "order": "0",
                "unEval": "false",
                "search": "",
                "from": "",
                "topicid": "0",
                "pages": str(page),
                "size": str(size),
            },
            timeout=timeout,
        )
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"mark list request failed: {exc}") from exc
    return response.text


def parse_submission_rows(html: str) -> tuple[list[dict[str, object]], int]:
    hidden = parse_hidden_inputs(html)
    total_pages = int(hidden.get("totalPage") or "1")
    rows: list[dict[str, object]] = []

    for match in re.finditer(r"<ul\b[^>]*class=['\"][^'\"]*dataBody_td[^'\"]*['\"][\s\S]*?</ul>", html, flags=re.I):
        block = match.group(0)
        open_tag = re.search(r"<ul\b[^>]*>", block, flags=re.I)
        attrs = html_attrs(open_tag.group(0)) if open_tag else {}
        answer_id = str(attrs.get("id") or "").strip()
        if not answer_id:
            continue

        name_match = re.search(r"<div\b[^>]*class=['\"][^'\"]*py_name[^'\"]*['\"][^>]*>(.*?)</div>", block, flags=re.I | re.S)
        name = compact_text(name_match.group(1)) if name_match else ""
        li_texts = [compact_text(item) for item in re.findall(r"<li\b[^>]*>(.*?)</li>", block, flags=re.I | re.S)]
        details = [item for item in li_texts if item and item not in {"批阅 打回"}]
        score_input = re.search(r"<input\b[^>]*class=['\"][^'\"]*scoreInput[^'\"]*['\"][^>]*>", block, flags=re.I | re.S)
        score_attrs = html_attrs(score_input.group(0)) if score_input else {}
        review_url_match = re.search(r"data=['\"]([^'\"]*?/mooc2-ans/work/library/review-work\?[^'\"]+)['\"]", block, flags=re.I)
        review_url = urljoin(MOOC2_BASE_URL, unescape(review_url_match.group(1))) if review_url_match else ""

        rows.append(
            {
                "answer_id": answer_id,
                "person_id": str(attrs.get("createid") or ""),
                "student_name": name,
                "student_no": details[1] if len(details) > 1 else "",
                "submitted_at": details[2] if len(details) > 2 else "",
                "submit_ip": details[3] if len(details) > 3 else "",
                "status": details[4] if len(details) > 4 else "",
                "reviewer": details[5] if len(details) > 5 else "",
                "score": score_attrs.get("value", ""),
                "score_data": score_attrs.get("data", ""),
                "review_url": review_url,
            }
        )
    return rows, total_pages


def fetch_submissions_for_homework(
    cookie_file: str,
    course: dict[str, object],
    clazz: dict[str, object],
    work_id: str,
    status: int,
    timeout: float,
) -> list[dict[str, object]]:
    session = request_session_from_cookie_file(cookie_file)
    size = 100
    first_html = fetch_mark_list_page(session, course, clazz, work_id, status, 1, size, timeout)
    rows, total_pages = parse_submission_rows(first_html)
    for page in range(2, total_pages + 1):
        html = fetch_mark_list_page(session, course, clazz, work_id, status, page, size, timeout)
        page_rows, _total_pages = parse_submission_rows(html)
        rows.extend(page_rows)

    for index, row in enumerate(rows, 1):
        row["index"] = index
    return rows


def find_submission(submissions: list[dict[str, object]], query: str) -> dict[str, object]:
    normalized = query.strip().lower()
    if not normalized:
        raise ToolError("submission index/answer_id/student name cannot be empty")

    for key in ("answer_id", "index", "student_no"):
        exact = [row for row in submissions if str(row.get(key, "")).strip().lower() == normalized]
        if len(exact) == 1:
            return exact[0]

    exact_name = [
        row
        for row in submissions
        if str(row.get("student_name", "")).strip().lower() == normalized
    ]
    if len(exact_name) == 1:
        return exact_name[0]
    if len(exact_name) > 1:
        names = ", ".join(str(row.get("student_name")) for row in exact_name)
        raise ToolError(f"multiple exact student matches: {names}")

    partial = [
        row
        for row in submissions
        if normalized in str(row.get("student_name", "")).strip().lower()
    ]
    if len(partial) == 1:
        return partial[0]
    if not partial:
        raise ToolError(f"submission not found: {query}")

    names = ", ".join(f"{row['index']}. {row['student_name']}" for row in partial)
    raise ToolError(f"multiple submission matches: {names}")


def parse_score_range(value: str) -> tuple[float, float]:
    numbers = re.findall(r"\d+(?:\.\d+)?", value)
    if len(numbers) != 2:
        raise ToolError("score range must contain two numbers, for example: 80-90")
    low = float(numbers[0])
    high = float(numbers[1])
    if low < 0 or high < 0:
        raise ToolError("score range cannot be negative")
    if low > high:
        raise ToolError("score range lower bound cannot exceed upper bound")
    return low, high


def random_score(
    low: float,
    high: float,
    decimals: int,
    distribution: str = "uniform",
    mean: float = None,
    stddev: float = None,
) -> str:
    if decimals not in (0, 1):
        raise ToolError("Chaoxing scores support at most one decimal place")
    scale = 10**decimals
    low_units = int(round(low * scale))
    high_units = int(round(high * scale))

    if distribution == "uniform":
        value_units = random.randint(low_units, high_units)
    elif distribution == "normal":
        normal_mean = (low + high) / 2 if mean is None else mean
        normal_stddev = (high - low) / 6 if stddev is None else stddev
        if normal_stddev <= 0:
            raise ToolError("normal distribution stddev must be greater than 0")
        value_units = None
        for _attempt in range(1000):
            candidate = random.gauss(normal_mean, normal_stddev)
            candidate_units = int(round(candidate * scale))
            if low_units <= candidate_units <= high_units:
                value_units = candidate_units
                break
        if value_units is None:
            value_units = max(low_units, min(high_units, int(round(normal_mean * scale))))
    else:
        raise ToolError("distribution must be one of: uniform, normal")

    value = value_units / scale
    return str(int(value)) if decimals == 0 else f"{value:.1f}"


def mark_homework_score(
    session: object,
    course: dict[str, object],
    clazz: dict[str, object],
    work_id: str,
    answer_id: str,
    score: str,
    timeout: float,
) -> dict[str, object]:
    try:
        response = session.post(
            MOOC2_BASE_URL
            + "/mooc2-ans/work/markscore?"
            + urlencode({"markScore": score, "markAnswerIds": answer_id, "markType": "0"}),
            data={
                "courseid": str(course.get("course_id") or ""),
                "clazzid": str(clazz.get("clazz_id") or ""),
                "cpi": str(course.get("cpi") or ""),
                "workid": work_id,
                "answerIds": answer_id,
                "type": "0",
                "score": score,
            },
            timeout=timeout,
        )
        response.raise_for_status()
        result = response.json()
    except Exception as exc:
        raise ToolError(f"mark score request failed for answer_id={answer_id}: {exc}") from exc
    return result


def default_grade_plan_dir() -> str:
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), DEFAULT_GRADE_PLAN_DIR)


def default_review_bundle_dir() -> str:
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), DEFAULT_REVIEW_BUNDLE_DIR)


def default_question_draft_dir() -> str:
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), DEFAULT_QUESTION_DRAFT_DIR)


def save_grade_plan(payload: dict[str, object], plan_dir: str) -> str:
    os.makedirs(plan_dir, exist_ok=True)
    work = payload.get("homework") if isinstance(payload.get("homework"), dict) else {}
    work_id = str(work.get("work_id") or "work")
    filename = f"grade_plan_{time.strftime('%Y%m%d_%H%M%S')}_{work_id}.json"
    path = os.path.join(plan_dir, filename)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)
        file.write("\n")
    os.chmod(path, 0o600)
    return path


def save_review_bundle(payload: dict[str, object], bundle_dir: str) -> str:
    os.makedirs(bundle_dir, exist_ok=True)
    work = payload.get("homework") if isinstance(payload.get("homework"), dict) else {}
    work_id = str(work.get("work_id") or "work")
    filename = f"review_bundle_{time.strftime('%Y%m%d_%H%M%S')}_{work_id}.json"
    path = os.path.join(bundle_dir, filename)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)
        file.write("\n")
    os.chmod(path, 0o600)
    return path


def save_question_draft(payload: dict[str, object], draft_dir: str, output: str = "") -> str:
    if output:
        path = os.path.abspath(output)
        parent = os.path.dirname(path)
        if parent:
            os.makedirs(parent, exist_ok=True)
    else:
        os.makedirs(draft_dir, exist_ok=True)
        title = str(payload.get("title") or "questions")
        filename = f"question_draft_{time.strftime('%Y%m%d_%H%M%S')}_{safe_filename(title)}.json"
        path = os.path.join(draft_dir, filename)
    with open(path, "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=2)
        file.write("\n")
    os.chmod(path, 0o600)
    return path


def safe_filename(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]+", "_", value).strip("._")
    return value[:120] or "asset"


def download_review_asset(session: object, url: str, asset_dir: str, prefix: str, timeout: float) -> str:
    os.makedirs(asset_dir, exist_ok=True)
    url = urljoin(MOOC2_BASE_URL, url)
    parsed = urlparse(url)
    ext = os.path.splitext(parsed.path)[1] or ".bin"
    filename = safe_filename(prefix) + ext
    path = os.path.join(asset_dir, filename)
    try:
        response = session.get(url, timeout=timeout)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"asset download failed: {url}: {exc}") from exc
    with open(path, "wb") as file:
        file.write(response.content)
    return path


def download_attachment(
    session: object,
    objectid: str,
    filename: str,
    asset_dir: str,
    timeout: float,
) -> str:
    if not objectid:
        raise ToolError("attachment objectid is empty")
    headers = {
        "Referer": MOOC2_BASE_URL + "/",
        "User-Agent": "Mozilla/5.0",
        "Accept": "application/json,text/plain,*/*",
    }
    try:
        status_response = session.get(
            MOOC2_BASE_URL + "/ananas/status/" + objectid,
            headers=headers,
            timeout=timeout,
        )
        status_response.raise_for_status()
        status = status_response.json()
    except Exception as exc:
        raise ToolError(f"attachment status request failed for objectid={objectid}: {exc}") from exc

    download_url = str(status.get("download") or "")
    if not download_url:
        raise ToolError(f"attachment has no download URL: objectid={objectid}")
    download_url = download_url.replace("http://", "https://", 1)

    os.makedirs(asset_dir, exist_ok=True)
    path = os.path.join(asset_dir, safe_filename(filename or objectid))
    try:
        response = session.get(download_url, headers=headers, timeout=timeout)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"attachment download failed for objectid={objectid}: {exc}") from exc
    with open(path, "wb") as file:
        file.write(response.content)
    return path


def extract_docx_text_fallback(path: str, max_chars: int) -> str:
    import zipfile
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
    root = ET.fromstring(document_xml)
    paragraphs = []
    for paragraph in root.iter():
        if not paragraph.tag.endswith("}p"):
            continue
        texts = [node.text for node in paragraph.iter() if node.tag.endswith("}t") and node.text]
        line = "".join(texts).strip()
        if line:
            paragraphs.append(line)
    return "\n".join(paragraphs)[:max_chars]


def extract_document_text(path: str, max_chars: int) -> str:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".docx", ".doc"):
            try:
                import docx

                document = docx.Document(path)
                parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
                return "\n".join(parts)[:max_chars]
            except Exception:
                return extract_docx_text_fallback(path, max_chars)
        if ext == ".pdf":
            import pdfplumber

            parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    if text.strip():
                        parts.append(text.strip())
            return "\n".join(parts)[:max_chars]
        if ext in (".txt", ".py", ".csv", ".md"):
            with open(path, "r", encoding="utf-8", errors="ignore") as file:
                return file.read(max_chars)
        if ext == ".zip":
            import zipfile

            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
            return "ZIP entries:\n" + "\n".join(names[:200])
    except Exception as exc:
        return f"(failed to extract {os.path.basename(path)}: {exc})"
    return ""


def normalize_question_types(value: str) -> list[str]:
    if not value.strip():
        raise ToolError("--types cannot be empty")
    types = []
    for raw_item in re.split(r"[,，/、\s]+", value.strip()):
        if not raw_item:
            continue
        question_type = QUESTION_TYPE_ALIASES.get(raw_item.strip().lower())
        if not question_type:
            raise ToolError(f"unsupported question type: {raw_item}")
        if question_type not in types:
            types.append(question_type)
    if not types:
        raise ToolError("--types cannot be empty")
    return types


def parse_csvish_list(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"[,，;；、\n]+", value or "") if item.strip()]


def clean_question_source_text(text: str, source_type: str) -> str:
    if source_type in ("md", "markdown", "text", "txt", "stdin"):
        text = re.sub(r"```[\s\S]*?```", " ", text)
        text = re.sub(r"!\[[^\]]*\]\([^)]+\)", " ", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        cleaned_lines = []
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                cleaned_lines.append("")
                continue
            line = re.sub(r"^#{1,6}\s*", "", line)
            line = re.sub(r"^[-*+]\s+", "", line)
            line = re.sub(r"^\d+[.)、]\s+", "", line)
            line = re.sub(r"^\|?[-: ]+\|[-|: ]+$", "", line)
            line = line.strip()
            if not line:
                continue
            if re.fullmatch(r"[A-Za-z0-9._/\- ]{1,40}", line):
                continue
            if len(line) <= 16 and not re.search(r"[。！？!?；;，,：:]", line):
                continue
            cleaned_lines.append(line)
        text = "\n".join(cleaned_lines)
    elif source_type in ("html", "htm"):
        text = compact_text(text)
    return compact_text(text)


def looks_like_cookie_payload(payload: object) -> bool:
    if not isinstance(payload, dict):
        return False
    cookies = payload.get("cookies")
    if not isinstance(cookies, list) or not cookies:
        return False
    return any(
        isinstance(cookie, dict) and {"name", "value", "domain"}.issubset(set(cookie.keys()))
        for cookie in cookies
    )


def collect_json_text(value: object, parts: list[str], key_path: str = "") -> None:
    sensitive_tokens = ("cookie", "password", "token", "secret", "authorization", "ticket")
    interesting_keys = (
        "title",
        "name",
        "chapter",
        "content",
        "text",
        "summary",
        "extracted_text",
        "question",
        "answer",
        "analysis",
        "description",
        "knowledge",
        "knowledge_point",
    )
    if any(token in key_path.lower() for token in sensitive_tokens):
        return

    if isinstance(value, dict):
        for key, child in value.items():
            child_path = f"{key_path}.{key}" if key_path else str(key)
            collect_json_text(child, parts, child_path)
        return

    if isinstance(value, list):
        for index, child in enumerate(value):
            collect_json_text(child, parts, f"{key_path}[{index}]")
        return

    if not isinstance(value, (str, int, float)):
        return

    text = compact_text(str(value))
    if not text:
        return
    key_name = key_path.rsplit(".", 1)[-1].lower()
    key_name = re.sub(r"\[\d+\]$", "", key_name)
    if key_name in interesting_keys or len(text) >= 20:
        parts.append(text)


def extract_json_source_text(path: str, max_chars: int) -> str:
    with open(path, "r", encoding="utf-8") as file:
        payload = json.load(file)
    if looks_like_cookie_payload(payload):
        raise ToolError("refusing to use a cookie file as question source")

    parts: list[str] = []
    if isinstance(payload, dict) and isinstance(payload.get("reviews"), list):
        course = payload.get("course") if isinstance(payload.get("course"), dict) else {}
        clazz = payload.get("selected_class") if isinstance(payload.get("selected_class"), dict) else {}
        homework = payload.get("homework") if isinstance(payload.get("homework"), dict) else {}
        for value in (
            course.get("course_name"),
            clazz.get("clazz_name"),
            homework.get("title"),
        ):
            if value:
                parts.append(str(value))
    collect_json_text(payload, parts)
    return "\n".join(dict.fromkeys(parts))[:max_chars]


def load_question_source(path: str, max_chars: int) -> dict[str, object]:
    if path == "-":
        text = clean_question_source_text(sys.stdin.read(max_chars), "stdin")
        return {
            "path": "-",
            "source_type": "stdin",
            "chars": len(text),
            "text": text,
        }

    if not os.path.exists(path):
        raise ToolError(f"source file not found: {path}")
    if os.path.isdir(path):
        raise ToolError(f"source path is a directory, not a file: {path}")

    ext = os.path.splitext(path)[1].lower()
    if re.search(r"cookie", os.path.basename(path), flags=re.I):
        raise ToolError("refusing to use a cookie file as question source")

    if ext == ".json":
        text = extract_json_source_text(path, max_chars)
    elif ext in (".docx", ".doc", ".pdf"):
        text = extract_document_text(path, max_chars)
    else:
        with open(path, "r", encoding="utf-8", errors="ignore") as file:
            raw = file.read(max_chars)
        text = raw.strip()

    source_type = ext.lstrip(".") or "text"
    text = clean_question_source_text(text, source_type)

    return {
        "path": os.path.abspath(path),
        "source_type": source_type,
        "chars": len(text),
        "text": text,
    }


def split_source_units(text: str) -> list[str]:
    text = compact_text(text)
    paragraphs = [part.strip() for part in re.split(r"\n{2,}", text) if part.strip()]
    units: list[str] = []
    for paragraph in paragraphs or [text]:
        chunks = re.split(r"(?<=[。！？!?；;：:])\s*", paragraph)
        buffer = ""
        for chunk in chunks:
            chunk = compact_text(chunk)
            if not chunk:
                continue
            if len(chunk) < 18 and buffer:
                buffer = f"{buffer}{chunk}"
                continue
            if buffer:
                units.append(buffer)
            buffer = chunk
        if buffer:
            units.append(buffer)

    cleaned = []
    seen = set()
    for unit in units:
        unit = re.sub(r"\s+", " ", unit).strip()
        if len(unit) < 12:
            continue
        if len(unit) > 260:
            pieces = re.split(r"(?<=[，,、])", unit)
            unit = "".join(pieces[:4]).strip() or unit[:260]
        key = unit[:120]
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(unit)
    return cleaned


def source_unit_score(unit: str) -> int:
    score = 0
    if 28 <= len(unit) <= 180:
        score += 4
    if re.search(r"(是|指|包括|包含|分为|采用|通过|用于|能够|可以|需要|核心|关键|因此|所以)", unit):
        score += 3
    if re.search(r"[\u4e00-\u9fff]", unit):
        score += 2
    if re.search(r"\d|%|[A-Za-z]{2,}", unit):
        score += 1
    return score


def select_question_units(text: str, count: int) -> list[str]:
    units = split_source_units(text)
    if not units:
        raise ToolError("source content is empty after text extraction")
    ranked = sorted(enumerate(units), key=lambda item: (-source_unit_score(item[1]), item[0]))
    selected = [unit for _index, unit in ranked[: max(count * 2, count)]]
    if len(selected) < count:
        selected = units
    return selected


def extract_question_topic(unit: str) -> str:
    text = compact_text(unit)
    text = re.sub(r"^[\d一二三四五六七八九十]+[、.．]\s*", "", text)
    for pattern in (
        r"(.{2,28}?)(?:是|指|包括|包含|分为|采用|通过|用于|能够|可以|需要|属于)",
        r"(?:关于|围绕|针对)(.{2,24}?)(?:，|,|。|；|;)",
    ):
        match = re.search(pattern, text)
        if match:
            topic = re.sub(r"^(因此|所以|其中|同时|此外|本节|本文|材料中)", "", match.group(1)).strip()
            topic = topic.strip(" ：:，,。；;“”\"'（）()[]【】")
            topic = re.sub(r"(不|未|没有)$", "", topic).strip()
            if 2 <= len(topic) <= 24:
                return topic
    words = re.findall(r"[A-Za-z][A-Za-z0-9_\-]{2,}|[\u4e00-\u9fff]{2,8}", text)
    stopwords = {"因此", "所以", "可以", "需要", "包括", "通过", "进行", "主要", "能够", "材料"}
    for word in words:
        if word not in stopwords:
            return word[:18]
    return text[:18]


def short_option_text(value: str, limit: int = 96) -> str:
    return short_text(value, limit).rstrip("。；;")


def choose_distractors(units: list[str], correct: str, topic: str, needed: int) -> list[str]:
    distractors = []
    topic = topic.strip()
    for unit in units:
        if unit == correct:
            continue
        if topic and topic in unit:
            continue
        option = short_option_text(unit)
        if option and option not in distractors:
            distractors.append(option)
        if len(distractors) >= needed:
            break
    fillers = [
        "只关注材料未展开的外部背景",
        "忽略材料中的限定条件",
        "把局部现象直接等同于整体结论",
    ]
    for filler in fillers:
        if len(distractors) >= needed:
            break
        distractors.append(filler)
    return distractors[:needed]


def option_rows(texts: list[str]) -> list[dict[str, str]]:
    labels = ["A", "B", "C", "D", "E", "F"]
    return [{"label": labels[index], "text": text} for index, text in enumerate(texts)]


def rotate_options(correct: str, distractors: list[str], index: int) -> tuple[list[dict[str, str]], str]:
    texts = [short_option_text(correct), *distractors[:3]]
    offset = index % len(texts)
    rotated = texts[offset:] + texts[:offset]
    options = option_rows(rotated)
    answer = next(option["label"] for option in options if option["text"] == short_option_text(correct))
    return options, answer


def mixed_difficulty(question_type: str, index: int) -> str:
    if question_type in ("true_false", "single_choice"):
        return "easy" if index % 3 else "medium"
    if question_type == "essay":
        return "hard"
    return "medium"


def build_question(
    index: int,
    question_type: str,
    unit: str,
    units: list[str],
    difficulty: str,
) -> dict[str, object]:
    topic = extract_question_topic(unit)
    difficulty_value = mixed_difficulty(question_type, index) if difficulty == "mixed" else difficulty
    score = DEFAULT_QUESTION_SCORES[question_type]
    base = {
        "id": f"Q{index:03d}",
        "type": question_type,
        "type_name": QUESTION_TYPE_NAMES[question_type],
        "difficulty": difficulty_value,
        "score": score,
        "source_evidence": unit,
        "draft_status": "needs_teacher_review",
    }

    if question_type == "single_choice":
        options, answer = rotate_options(unit, choose_distractors(units, unit, topic, 3), index)
        return {
            **base,
            "stem": f"根据材料，关于“{topic}”，下列哪项表述最符合原文？",
            "options": options,
            "answer": answer,
            "analysis": f"原文依据：{unit}",
        }

    if question_type == "multiple_choice":
        related = [unit]
        for candidate in units:
            if candidate != unit and len(related) < 2 and topic and topic in candidate:
                related.append(candidate)
        for candidate in units:
            if candidate != unit and candidate not in related and len(related) < 2:
                related.append(candidate)
        distractors = choose_distractors(units, unit, topic, 4 - len(related))
        texts = [short_option_text(item) for item in related] + distractors
        options = option_rows(texts[:4])
        answers = [option["label"] for option in options if option["text"] in {short_option_text(item) for item in related}]
        return {
            **base,
            "stem": f"根据材料，下列哪些表述与“{topic}”直接相关？",
            "options": options,
            "answer": answers,
            "analysis": "正确选项均可在材料对应语句中找到依据。",
        }

    if question_type == "true_false":
        if index % 2:
            statement = short_option_text(unit, 140)
            answer = "正确"
            analysis = f"该表述来自材料：{unit}"
        else:
            distractor = choose_distractors(units, unit, topic, 1)[0]
            statement = f"材料认为“{topic}”的核心内容是：{distractor}。"
            answer = "错误"
            analysis = f"该表述混淆了材料内容。关于“{topic}”的依据是：{unit}"
        return {
            **base,
            "stem": statement,
            "answer": answer,
            "analysis": analysis,
        }

    if question_type == "short_answer":
        return {
            **base,
            "stem": f"简述材料中“{topic}”的核心内容。",
            "reference_answer": unit,
            "answer": unit,
            "analysis": f"答题要点应覆盖材料中的关键信息：{unit}",
        }

    if question_type == "essay":
        unit_index = units.index(unit) if unit in units else 0
        evidence_units = units[unit_index : unit_index + 3] or [unit]
        return {
            **base,
            "stem": f"结合材料，分析“{topic}”的作用、实现思路或教学启示。",
            "reference_answer": "\n".join(evidence_units),
            "scoring_points": [
                "能准确概括材料中的核心概念或过程。",
                "能结合材料说明关键依据，而不是只给泛泛结论。",
                "能展开分析作用、限制、应用场景或教学启示。",
            ],
            "analysis": "论述题答案应围绕材料证据展开，可按课堂要求补充案例。",
        }

    raise ToolError(f"unsupported question type: {question_type}")


def build_question_draft_payload(args: argparse.Namespace) -> dict[str, object]:
    if args.count <= 0:
        raise ToolError("--count must be greater than 0")
    if args.max_source_chars <= 0:
        raise ToolError("--max-source-chars must be greater than 0")

    question_types = normalize_question_types(args.types)
    per_source_limit = max(1000, args.max_source_chars // max(len(args.sources), 1))
    sources = [load_question_source(path, per_source_limit) for path in args.sources]
    combined_text = compact_text("\n\n".join(str(source.get("text") or "") for source in sources))
    if not combined_text:
        raise ToolError("no usable source text found")
    combined_text = combined_text[: args.max_source_chars]

    units = select_question_units(combined_text, args.count)
    questions = []
    for index in range(1, args.count + 1):
        question_type = question_types[(index - 1) % len(question_types)]
        unit = units[(index - 1) % len(units)]
        questions.append(build_question(index, question_type, unit, units, args.difficulty))

    title = args.title.strip() or "内容出题草稿"
    return {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "title": title,
        "metadata": {
            "course": args.course,
            "chapter": args.chapter,
            "knowledge_points": parse_csvish_list(args.knowledge_points),
            "question_types": question_types,
            "difficulty": args.difficulty,
            "requested_count": args.count,
        },
        "source_files": [
            {
                "path": source.get("path"),
                "source_type": source.get("source_type"),
                "chars": source.get("chars"),
                "text_excerpt": short_text(source.get("text"), 500),
            }
            for source in sources
        ],
        "source_excerpt": combined_text[:3000],
        "generation_policy": {
            "mode": "offline_draft",
            "published_to_chaoxing": False,
            "teacher_review_required": True,
            "note": "题目由本地材料抽取和模板生成，发布或导入前需要教师审题、改写干扰项并确认答案。",
        },
        "questions": questions,
    }


def render_question_markdown(payload: dict[str, object]) -> str:
    lines = [f"# {payload.get('title') or '内容出题草稿'}", ""]
    metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
    if metadata.get("course"):
        lines.append(f"- 课程：{metadata.get('course')}")
    if metadata.get("chapter"):
        lines.append(f"- 章节：{metadata.get('chapter')}")
    if metadata.get("knowledge_points"):
        lines.append("- 知识点：" + "、".join(str(item) for item in metadata.get("knowledge_points", [])))
    lines.append("- 状态：草稿，发布前需要教师复核")
    lines.append("")

    questions = payload.get("questions") if isinstance(payload.get("questions"), list) else []
    for index, question in enumerate(questions, 1):
        if not isinstance(question, dict):
            continue
        lines.append(f"## {index}. [{question.get('type_name')}] {question.get('stem')}")
        options = question.get("options")
        if isinstance(options, list):
            for option in options:
                if isinstance(option, dict):
                    lines.append(f"- {option.get('label')}. {option.get('text')}")
        answer = question.get("answer")
        if isinstance(answer, list):
            answer_text = "、".join(str(item) for item in answer)
        else:
            answer_text = str(answer or question.get("reference_answer") or "")
        if answer_text:
            lines.append(f"答案：{answer_text}")
        if question.get("analysis"):
            lines.append(f"解析：{question.get('analysis')}")
        if question.get("source_evidence"):
            lines.append(f"依据：{question.get('source_evidence')}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def print_question_draft_summary(payload: dict[str, object], draft_file: str) -> None:
    questions = payload.get("questions") if isinstance(payload.get("questions"), list) else []
    print(f"draft_file: {draft_file}")
    print(f"title: {payload.get('title')}")
    print(f"questions: {len(questions)}")
    print("draft_status: not published; teacher review required before importing to Chaoxing")
    for question in questions:
        if not isinstance(question, dict):
            continue
        answer = question.get("answer")
        if isinstance(answer, list):
            answer_text = ",".join(str(item) for item in answer)
        else:
            answer_text = str(answer or "")
        answer_suffix = f" answer={answer_text}" if answer_text else ""
        print(
            f"  {question.get('id')}. [{question.get('type_name')}] "
            f"{short_text(question.get('stem'), 90)}{answer_suffix}"
        )


def command_exam_question_draft(args: argparse.Namespace) -> int:
    payload = build_question_draft_payload(args)
    draft_file = save_question_draft(payload, args.draft_dir, args.output)
    markdown_file = ""
    markdown = ""
    if args.markdown_output or args.format == "markdown":
        markdown = render_question_markdown(payload)
    if args.markdown_output:
        markdown_file = os.path.abspath(args.markdown_output)
        parent = os.path.dirname(markdown_file)
        if parent:
            os.makedirs(parent, exist_ok=True)
        with open(markdown_file, "w", encoding="utf-8") as file:
            file.write(markdown)
        os.chmod(markdown_file, 0o600)

    if args.format == "json":
        print(
            json.dumps(
                {
                    "draft_file": draft_file,
                    "markdown_file": markdown_file,
                    **payload,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    elif args.format == "markdown":
        print(markdown.rstrip())
    else:
        print_question_draft_summary(payload, draft_file)
        if markdown_file:
            print(f"markdown_file: {markdown_file}")
    return 0


def is_answer_image_url(url: str) -> bool:
    if not url:
        return False
    ignored = (
        "/mooc2-ans/module/work/images/",
        "/mooc2-ans/images/",
        "popClose",
    )
    return not any(token in url for token in ignored)


def parse_attachment_iframes(fragment: str) -> list[dict[str, str]]:
    attachments = []
    for iframe in re.findall(r"<iframe\b[^>]*class=['\"][^'\"]*attach-iframe[^'\"]*['\"][^>]*>", fragment, flags=re.I | re.S):
        attrs = html_attrs(iframe)
        attachments.append(
            {
                "filename": attrs.get("filename", ""),
                "filetype": attrs.get("filetype", ""),
                "objectid": attrs.get("objectid", ""),
            }
        )
    return attachments


def load_grade_plan(path: str) -> dict[str, object]:
    if not os.path.exists(path):
        raise ToolError(f"grade plan not found: {path}")
    with open(path, "r", encoding="utf-8") as file:
        payload = json.load(file)
    if not isinstance(payload, dict):
        raise ToolError(f"invalid grade plan: {path}")
    plan = payload.get("plan")
    if not isinstance(plan, list) or not plan:
        raise ToolError(f"grade plan has no score rows: {path}")
    return payload


def load_review_bundle(path: str) -> dict[str, object]:
    if not os.path.exists(path):
        raise ToolError(f"review bundle not found: {path}")
    with open(path, "r", encoding="utf-8") as file:
        payload = json.load(file)
    if not isinstance(payload, dict):
        raise ToolError(f"invalid review bundle: {path}")
    reviews = payload.get("reviews")
    if not isinstance(reviews, list) or not reviews:
        raise ToolError(f"review bundle has no reviews: {path}")
    return payload


def load_review_scores(path: str) -> list[dict[str, object]]:
    if not os.path.exists(path):
        raise ToolError(f"score review file not found: {path}")
    with open(path, "r", encoding="utf-8") as file:
        payload = json.load(file)
    if isinstance(payload, dict) and isinstance(payload.get("scores"), list):
        rows = payload["scores"]
    elif isinstance(payload, list):
        rows = payload
    elif isinstance(payload, dict):
        rows = []
        for key, value in payload.items():
            if isinstance(value, dict):
                row = dict(value)
                row.setdefault("answer_id", key)
            else:
                row = {"answer_id": key, "score": value}
            rows.append(row)
    else:
        raise ToolError(f"invalid score review file: {path}")
    if not rows or not all(isinstance(row, dict) for row in rows):
        raise ToolError(f"score review file has no score rows: {path}")
    return rows


def score_to_float(value: object, label: str) -> float:
    try:
        return float(str(value).strip())
    except Exception as exc:
        raise ToolError(f"invalid score for {label}: {value}") from exc


def format_score(value: float, decimals: int) -> str:
    if decimals not in (0, 1):
        raise ToolError("Chaoxing scores support at most one decimal place")
    if decimals == 0:
        return str(int(round(value)))
    return f"{round(value, 1):.1f}"


def score_equal(left: object, right: object) -> bool:
    try:
        return abs(float(str(left).strip()) - float(str(right).strip())) < 0.05
    except Exception:
        return str(left).strip() == str(right).strip()


def score_row_key(row: dict[str, object]) -> str:
    for key in ("answer_id", "student_no", "student_name"):
        value = str(row.get(key) or "").strip()
        if value:
            return value
    return ""


def index_score_rows(rows: list[dict[str, object]]) -> dict[str, dict[str, object]]:
    indexed: dict[str, dict[str, object]] = {}
    for row in rows:
        keys = [str(row.get(key) or "").strip() for key in ("answer_id", "student_no", "student_name")]
        keys = [key for key in keys if key]
        if not keys:
            raise ToolError("every score row needs answer_id, student_no, or student_name")
        for key in keys:
            if key in indexed:
                raise ToolError(f"duplicate score row key: {key}")
            indexed[key] = row
    return indexed


def bundle_full_score(bundle: dict[str, object]) -> float:
    reviews = bundle.get("reviews") if isinstance(bundle.get("reviews"), list) else []
    for item in reviews:
        if not isinstance(item, dict):
            continue
        summary = item.get("summary") if isinstance(item.get("summary"), dict) else {}
        value = str(summary.get("full_score") or "").strip()
        if value:
            return score_to_float(value, "full_score")
    return 100.0


def build_review_grade_plan(
    bundle: dict[str, object],
    score_rows: list[dict[str, object]],
    score_range: tuple[float, float] | None,
    floor: float,
    decimals: int,
    require_reasons: bool,
    allow_partial: bool,
    source_bundle: str,
) -> dict[str, object]:
    course = bundle.get("course")
    clazz = bundle.get("selected_class")
    homework = bundle.get("homework")
    reviews = bundle.get("reviews")
    if not isinstance(course, dict) or not isinstance(clazz, dict) or not isinstance(homework, dict):
        raise ToolError("review bundle is missing course/class/homework metadata")
    if not isinstance(reviews, list):
        raise ToolError("review bundle has invalid reviews")

    full_score = bundle_full_score(bundle)
    low = 0.0 if score_range is None else score_range[0]
    high = full_score if score_range is None else score_range[1]
    if high > full_score:
        raise ToolError(f"score range upper bound {high:g} exceeds homework full score {full_score:g}")

    indexed_scores = index_score_rows(score_rows)
    plan = []
    matched_keys = set()
    missing = []
    for review in reviews:
        if not isinstance(review, dict):
            continue
        submission = review.get("submission") if isinstance(review.get("submission"), dict) else {}
        candidates = [
            str(submission.get("answer_id") or "").strip(),
            str(submission.get("student_no") or "").strip(),
            str(submission.get("student_name") or "").strip(),
        ]
        score_row = next((indexed_scores[key] for key in candidates if key and key in indexed_scores), None)
        if score_row is None:
            missing.append(str(submission.get("student_name") or submission.get("answer_id") or "unknown"))
            continue
        matched_keys.add(score_row_key(score_row))

        label = str(submission.get("student_name") or submission.get("answer_id") or "unknown")
        raw_score = score_row.get("score")
        if raw_score is None:
            raise ToolError(f"missing score for {label}")
        score = score_to_float(raw_score, label)
        if score < low or score > high:
            raise ToolError(f"score for {label} is outside allowed range {low:g}-{high:g}: {score:g}")

        reason = compact_text(str(score_row.get("reason") or ""))
        evidence = compact_text(str(score_row.get("evidence") or ""))
        content_summary = compact_text(str(score_row.get("content_summary") or ""))
        if require_reasons and not reason:
            raise ToolError(f"missing grading reason for {label}")
        if score < floor and len(reason) < 20:
            raise ToolError(f"score below {floor:g} for {label} needs a concrete reason")

        plan.append(
            {
                "index": submission.get("index"),
                "answer_id": submission.get("answer_id"),
                "student_name": submission.get("student_name"),
                "student_no": submission.get("student_no"),
                "score": format_score(score, decimals),
                "reason": reason,
                "evidence": evidence,
                "content_summary": content_summary,
            }
        )

    if missing and not allow_partial:
        raise ToolError("missing scores for: " + ", ".join(missing))
    if not plan:
        raise ToolError("no score rows matched the review bundle")

    matched_answer_ids = {str(item.get("answer_id") or "") for item in plan}
    extras = []
    for row in score_rows:
        keys = {str(row.get(key) or "").strip() for key in ("answer_id", "student_no", "student_name")}
        keys.discard("")
        if not any(
            key in matched_answer_ids
            or any(key == str(item.get("student_no") or "") or key == str(item.get("student_name") or "") for item in plan)
            for key in keys
        ):
            extras.append(score_row_key(row))
    if extras and not allow_partial:
        raise ToolError("score rows not found in review bundle: " + ", ".join(extras))

    return {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "course": course,
        "selected_class": clazz,
        "homework": homework,
        "work_info": {"score": full_score},
        "source_bundle": source_bundle,
        "score_range": {
            "min": low,
            "max": high,
            "decimals": decimals,
            "source": "content_review",
            "floor": floor,
        },
        "grading_policy": bundle.get("grading_policy", {}),
        "plan": plan,
    }


def verify_submitted_scores(
    cookie_file: str,
    course: dict[str, object],
    clazz: dict[str, object],
    homework: dict[str, object],
    plan: list[dict[str, object]],
    timeout: float,
) -> dict[str, object]:
    rows = fetch_submissions_for_homework(
        cookie_file,
        course,
        clazz,
        str(homework.get("work_id") or ""),
        0,
        timeout,
    )
    by_answer_id = {str(row.get("answer_id") or ""): row for row in rows}
    matched = []
    mismatches = []
    missing = []
    for item in plan:
        answer_id = str(item.get("answer_id") or "")
        row = by_answer_id.get(answer_id)
        if row is None:
            missing.append(item)
            continue
        expected = item.get("score")
        actual = row.get("score")
        result = {
            "answer_id": answer_id,
            "student_name": item.get("student_name"),
            "expected": expected,
            "actual": actual,
        }
        if score_equal(expected, actual):
            matched.append(result)
        else:
            mismatches.append(result)
    return {
        "expected": len(plan),
        "matched": len(matched),
        "mismatch": len(mismatches),
        "missing": len(missing),
        "matched_rows": matched,
        "mismatches": mismatches,
        "missing_rows": missing,
    }


def short_text(value: object, limit: int = 160) -> str:
    text = compact_text(str(value or ""))
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def print_grade_plan_summary(payload: dict[str, object]) -> None:
    course = payload.get("course") if isinstance(payload.get("course"), dict) else {}
    clazz = payload.get("selected_class") if isinstance(payload.get("selected_class"), dict) else {}
    work = payload.get("homework") if isinstance(payload.get("homework"), dict) else {}
    score_range = payload.get("score_range") if isinstance(payload.get("score_range"), dict) else {}
    plan = payload.get("plan") if isinstance(payload.get("plan"), list) else []

    print(f"course: {course.get('course_name')} ({course.get('course_id')})")
    print(f"selected_class: {clazz.get('clazz_name')} (clazz_id={clazz.get('clazz_id')})")
    print(f"homework: {work.get('title')} (work_id={work.get('work_id')})")
    print(
        f"score_range: {score_range.get('min')}-{score_range.get('max')}, "
        f"decimals={score_range.get('decimals')}"
    )
    print("final_score_summary:")
    for item in plan:
        if not isinstance(item, dict):
            continue
        reason = short_text(item.get("reason", ""))
        reason_suffix = f" | reason: {reason}" if reason else ""
        print(
            f"  {item.get('index')}. {item.get('student_name')} "
            f"(answer_id={item.get('answer_id')}, student_no={item.get('student_no')}) "
            f"-> {item.get('score')}{reason_suffix}"
        )


def extract_review_summary(html: str, max_chars: int) -> dict[str, object]:
    hidden = parse_hidden_inputs(html)
    answer_fragments = re.findall(
        r"<dd\b[^>]*class=['\"][^'\"]*stuAnswerWords[^'\"]*['\"][^>]*>([\s\S]*?)</dd>",
        html,
        flags=re.I,
    )
    student_answers = []
    answer_attachments = []
    for fragment in answer_fragments:
        images = [
            urljoin(MOOC2_BASE_URL, unescape(src))
            for src in re.findall(r"<img\b[^>]*src=['\"]([^'\"]+)['\"]", fragment, flags=re.I)
            if is_answer_image_url(src)
        ]
        attachments = parse_attachment_iframes(fragment)
        answer_attachments.extend(attachments)
        student_answers.append(
            {
                "text": compact_text(fragment),
                "images": images,
                "attachments": attachments,
            }
        )
    answer_images = []
    for answer in student_answers:
        answer_images.extend(answer["images"])
    attachments = parse_attachment_iframes(html)

    plain = compact_text(html)
    return {
        "work_answer_id": hidden.get("workAnswerId", ""),
        "work_id": hidden.get("workId", ""),
        "full_score": hidden.get("fullScore", ""),
        "student_answers": student_answers,
        "answer_images": answer_images,
        "student_answer_attachments": answer_attachments,
        "attachments": attachments,
        "text_excerpt": plain[:max_chars],
    }


def print_teaching_courses(courses: list[dict[str, object]]) -> None:
    for course in courses:
        details = " - ".join(
            str(value)
            for value in (course.get("term"), course.get("teacher"))
            if value
        )
        suffix = f" - {details}" if details else ""
        print(
            f"{course['index']}. {course['course_name']}{suffix} "
            f"(course_id={course['course_id']}, classes={course['class_count']})"
        )


def print_course_classes(course: dict[str, object]) -> None:
    print(f"course: {course['course_name']} ({course['course_id']})")
    if course.get("current_user"):
        print(f"current_user: {course['current_user']}")
    if course.get("term"):
        print(f"term: {course['term']}")
    if course.get("teacher"):
        print(f"teacher: {course['teacher']}")
    print("classes:")
    for clazz in course.get("classes", []):
        count = clazz.get("student_count")
        count_text = f", students={count}" if count is not None else ""
        marker = "[用户] " if clazz.get("is_user_class") else ""
        print(f"  {clazz['index']}. {marker}{clazz['clazz_name']} (clazz_id={clazz['clazz_id']}{count_text})")


def print_task_options(course: dict[str, object], clazz: dict[str, object]) -> None:
    marker = "[用户] " if clazz.get("is_user_class") else ""
    count = clazz.get("student_count")
    count_text = f", students={count}" if count is not None else ""

    print(f"course: {course['course_name']} ({course['course_id']})")
    if course.get("term"):
        print(f"term: {course['term']}")
    print(f"selected_class: {marker}{clazz['clazz_name']} (clazz_id={clazz['clazz_id']}{count_text})")
    print("tasks:")
    for option in task_options():
        print(
            f"  {option['index']}. {option['name']} "
            f"(task={option['task']}, pageHeader={option['page_header']})"
        )


def print_homework_items(course: dict[str, object], clazz: dict[str, object], items: list[dict[str, object]]) -> None:
    print(f"course: {course['course_name']} ({course['course_id']})")
    print(f"selected_class: {clazz['clazz_name']} (clazz_id={clazz['clazz_id']})")
    print("homeworks:")
    if not items:
        print("  (none)")
        return
    for item in items:
        print(
            f"  {item['index']}. {item['title']} "
            f"(work_id={item['work_id']}, pending={item['pending_count']}, "
            f"submitted={item['submitted_count']}, unsubmitted={item['unsubmitted_count']})"
        )


def print_homework_submissions(
    course: dict[str, object],
    clazz: dict[str, object],
    work: dict[str, object],
    submissions: list[dict[str, object]],
    work_info: dict[str, object],
) -> None:
    print(f"course: {course['course_name']} ({course['course_id']})")
    print(f"selected_class: {clazz['clazz_name']} (clazz_id={clazz['clazz_id']})")
    print(f"homework: {work['title']} (work_id={work['work_id']}, full_score={work_info.get('score', '')})")
    print("pending_submissions:")
    if not submissions:
        print("  (none)")
        return
    for row in submissions:
        print(
            f"  {row['index']}. {row['student_name']} "
            f"(answer_id={row['answer_id']}, student_no={row['student_no']}, "
            f"submitted_at={row['submitted_at']}, score={row['score']})"
        )


def command_courses(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    courses = teaching_courses_from_backclazz(fetch_backclazz_data(args.cookie_file, args.timeout))
    if args.format == "json":
        print(json.dumps({"courses": courses}, ensure_ascii=False, indent=2))
    else:
        print_teaching_courses(courses)
    return 0


def command_course_classes(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    courses = teaching_courses_from_backclazz(fetch_backclazz_data(args.cookie_file, args.timeout))
    course = find_teaching_course(courses, args.course)
    owner_name = args.owner_name or detect_current_user_name(args.cookie_file, args.timeout)
    prioritize_user_classes(course, owner_name)
    if args.format == "json":
        print(json.dumps({"course": course}, ensure_ascii=False, indent=2))
    else:
        print_course_classes(course)
    return 0


def command_task_options(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    course, clazz = resolve_course_and_class(args)
    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "tasks": task_options(),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print_task_options(course, clazz)
    return 0


def command_course_task(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    task_kind = normalize_task_kind(args.task)
    task = TASK_DEFINITIONS[task_kind]
    course, clazz = resolve_course_and_class(args)
    context = fetch_teacher_task_context(
        args.cookie_file,
        course,
        clazz,
        int(task["page_header"]),
        args.timeout,
    )
    task_url = build_task_content_url(context, task_kind, clazz)
    entry_url = build_teacher_entry_url(course, clazz, int(task["page_header"]))

    if args.open:
        ensure_agent_browser()
        browser_args = base_args(args.session, args.headed)
        injected = inject_cookie_file_into_browser(browser_args, args.cookie_file)
        if injected:
            print(f"injected_cookies: {injected}")
        run_browser([*browser_args, "open", args.check_url])
        run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
        run_browser([*browser_args, "open", entry_url])
        run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
        run_browser([*browser_args, "open", task_url])

    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "task": {
                        "task": task_kind,
                        "name": task["name"],
                        "module": task["module"],
                        "page_header": task["page_header"],
                    },
                    "urls": {
                        "course_entry_url": entry_url,
                        "task_url": task_url,
                    },
                    "opened": bool(args.open),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(f"course: {course['course_name']} ({course['course_id']})")
        print(f"selected_class: {clazz['clazz_name']} (clazz_id={clazz['clazz_id']})")
        print(f"task: {task['name']} (task={task_kind}, pageHeader={task['page_header']})")
        print(f"course_entry_url: {entry_url}")
        print(f"task_url: {task_url}")
        if args.open:
            print("opened: true")
    return 0


def command_homework_ungraded(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    course, clazz = resolve_course_and_class(args)
    items = fetch_homework_items_for_class(args.cookie_file, course, clazz, args.timeout)
    selected = items if args.all else ungraded_homework_items(items)

    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "homeworks": selected,
                    "only_ungraded": not args.all,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print_homework_items(course, clazz, selected)
    return 0


def command_homework_submissions(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    course, clazz = resolve_course_and_class(args)
    items = fetch_homework_items_for_class(args.cookie_file, course, clazz, args.timeout)
    candidates = items if args.all_homeworks else ungraded_homework_items(items)
    work = find_homework_item(candidates, args.work)

    session = request_session_from_cookie_file(args.cookie_file)
    work_info = fetch_work_info(session, course, clazz, str(work["work_id"]), args.timeout)
    submissions = fetch_submissions_for_homework(
        args.cookie_file,
        course,
        clazz,
        str(work["work_id"]),
        args.status,
        args.timeout,
    )

    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "homework": work,
                    "work_info": work_info,
                    "submissions": submissions,
                    "status": args.status,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print_homework_submissions(course, clazz, work, submissions, work_info)
    return 0


def command_homework_review(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    course, clazz = resolve_course_and_class(args)
    items = fetch_homework_items_for_class(args.cookie_file, course, clazz, args.timeout)
    candidates = items if args.all_homeworks else ungraded_homework_items(items)
    work = find_homework_item(candidates, args.work)
    submissions = fetch_submissions_for_homework(
        args.cookie_file,
        course,
        clazz,
        str(work["work_id"]),
        args.status,
        args.timeout,
    )
    if not submissions:
        raise ToolError(f"no submissions found for homework: {work['title']}")
    submission = find_submission(submissions, args.answer) if args.answer else submissions[0]
    review_url = str(submission.get("review_url") or "")
    if not review_url:
        raise ToolError(f"submission has no review URL: {submission.get('answer_id')}")

    session = request_session_from_cookie_file(args.cookie_file)
    try:
        response = session.get(review_url, timeout=args.timeout, allow_redirects=True)
        response.raise_for_status()
    except Exception as exc:
        raise ToolError(f"review page request failed: {exc}") from exc
    summary = extract_review_summary(response.text, args.max_chars)

    if args.open:
        ensure_agent_browser()
        browser_args = base_args(args.session, args.headed)
        injected = inject_cookie_file_into_browser(browser_args, args.cookie_file)
        if injected:
            print(f"injected_cookies: {injected}")
        run_browser([*browser_args, "open", review_url])
        run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)

    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "homework": work,
                    "submission": submission,
                    "review": summary,
                    "review_url": review_url,
                    "opened": bool(args.open),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(f"homework: {work['title']} (work_id={work['work_id']})")
        print(
            f"submission: {submission['student_name']} "
            f"(answer_id={submission['answer_id']}, student_no={submission['student_no']})"
        )
        print(f"full_score: {summary.get('full_score')}")
        print(f"answer_images: {len(summary.get('answer_images', []))}")
        for index, image_url in enumerate(summary.get("answer_images", []), 1):
            print(f"  image_{index}: {image_url}")
        print(f"attachments: {len(summary.get('attachments', []))}")
        for index, attachment in enumerate(summary.get("attachments", []), 1):
            print(
                f"  attachment_{index}: {attachment.get('filename')} "
                f"({attachment.get('filetype')}, objectid={attachment.get('objectid')})"
            )
        print("text_excerpt:")
        print(summary.get("text_excerpt", ""))
        if args.open:
            print("opened: true")
    return 0


def command_homework_review_bundle(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    course, clazz = resolve_course_and_class(args)
    items = fetch_homework_items_for_class(args.cookie_file, course, clazz, args.timeout)
    candidates = items if args.all_homeworks else ungraded_homework_items(items)
    work = find_homework_item(candidates, args.work)
    submissions = fetch_submissions_for_homework(
        args.cookie_file,
        course,
        clazz,
        str(work["work_id"]),
        args.status,
        args.timeout,
    )
    if args.limit is not None and args.limit <= 0:
        raise ToolError("--limit must be greater than 0")
    if args.limit is not None:
        submissions = submissions[: args.limit]
    if not submissions:
        raise ToolError(f"no submissions found for homework: {work['title']}")

    session = request_session_from_cookie_file(args.cookie_file)
    reviews = []
    asset_dir = ""
    if args.download_assets:
        asset_dir = os.path.join(
            args.bundle_dir,
            "assets_" + time.strftime("%Y%m%d_%H%M%S") + "_" + str(work["work_id"]),
        )
    for row in submissions:
        review_url = str(row.get("review_url") or "")
        if not review_url:
            continue
        try:
            response = session.get(review_url, timeout=args.timeout, allow_redirects=True)
            response.raise_for_status()
        except Exception as exc:
            raise ToolError(f"review page request failed for answer_id={row.get('answer_id')}: {exc}") from exc
        summary = extract_review_summary(response.text, args.max_chars)
        downloaded_assets = []
        if args.download_assets:
            for index, image_url in enumerate(summary.get("answer_images", []), 1):
                local_path = download_review_asset(
                    session,
                    str(image_url),
                    asset_dir,
                    f"{row.get('index')}_{row.get('student_name')}_answer_{index}",
                    args.timeout,
                )
                downloaded_assets.append({"url": image_url, "path": local_path})
            for attachment in summary.get("student_answer_attachments", []):
                if not isinstance(attachment, dict):
                    continue
                filename = (
                    f"{row.get('index')}_{row.get('student_name')}_"
                    f"{attachment.get('filename') or attachment.get('objectid') or 'attachment'}"
                )
                local_path = download_attachment(
                    session,
                    str(attachment.get("objectid") or ""),
                    filename,
                    asset_dir,
                    args.timeout,
                )
                extracted_text = extract_document_text(local_path, args.max_chars)
                downloaded_assets.append(
                    {
                        "filename": attachment.get("filename", ""),
                        "filetype": attachment.get("filetype", ""),
                        "objectid": attachment.get("objectid", ""),
                        "path": local_path,
                        "extracted_text": extracted_text,
                    }
                )
        reviews.append(
            {
                "submission": row,
                "review_url": review_url,
                "summary": summary,
                "downloaded_assets": downloaded_assets,
            }
        )

    payload = {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "course": course_summary(course),
        "selected_class": class_summary(clazz),
        "homework": work,
        "status": args.status,
        "grading_policy": {
            "default_floor": 70,
            "below_70_requires_reason": True,
            "instruction": "建议分数尽量不低于70；低于70必须说明作业内容严重不合适的具体理由。",
        },
        "reviews": reviews,
    }
    bundle_file = save_review_bundle(payload, args.bundle_dir)

    if args.format == "json":
        print(json.dumps({"bundle_file": bundle_file, "asset_dir": asset_dir, **payload}, ensure_ascii=False, indent=2))
    else:
        print(f"bundle_file: {bundle_file}")
        if asset_dir:
            print(f"asset_dir: {asset_dir}")
        print(f"course: {course['course_name']} ({course['course_id']})")
        print(f"selected_class: {clazz['clazz_name']} (clazz_id={clazz['clazz_id']})")
        print(f"homework: {work['title']} (work_id={work['work_id']})")
        print(f"reviews: {len(reviews)}")
        print("grading_policy: default >=70; below 70 requires a concrete reason")
        for item in reviews:
            row = item["submission"]
            summary = item["summary"]
            print(
                f"  {row['index']}. {row['student_name']} "
                f"(answer_id={row['answer_id']}): "
                f"answer_images={len(summary.get('answer_images', []))}, "
                f"student_attachments={len(summary.get('student_answer_attachments', []))}"
            )
    return 0


def command_homework_grade_reviewed(args: argparse.Namespace) -> int:
    review_bundle = load_review_bundle(args.review_bundle)
    score_rows = load_review_scores(args.scores_file)
    score_range = parse_score_range(args.score_range) if args.score_range else None
    if args.floor < 0:
        raise ToolError("--floor cannot be negative")

    plan_payload = build_review_grade_plan(
        review_bundle,
        score_rows,
        score_range,
        args.floor,
        args.decimals,
        not args.allow_missing_reasons,
        args.allow_partial,
        args.review_bundle,
    )
    plan_file = save_grade_plan(plan_payload, args.plan_dir)

    if args.format == "json":
        print(
            json.dumps(
                {
                    "plan_file": plan_file,
                    "plan": plan_payload,
                    "dry_run": True,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(f"plan_file: {plan_file}")
        print_grade_plan_summary(plan_payload)
        print("not submitted; ask the user to confirm, then run homework-submit-plan with --commit")
    return 0


def command_homework_grade_random(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    low, high = parse_score_range(args.score_range)
    if args.seed is not None:
        random.seed(args.seed)

    course, clazz = resolve_course_and_class(args)
    items = fetch_homework_items_for_class(args.cookie_file, course, clazz, args.timeout)
    candidates = items if args.all_homeworks else ungraded_homework_items(items)
    work = find_homework_item(candidates, args.work)

    session = request_session_from_cookie_file(args.cookie_file)
    work_info = fetch_work_info(session, course, clazz, str(work["work_id"]), args.timeout)
    full_score = float(work_info.get("score") or 0)
    if high > full_score:
        raise ToolError(f"score range upper bound {high:g} exceeds homework full score {full_score:g}")
    if args.mean is not None and not (low <= args.mean <= high):
        raise ToolError("--mean must be inside the score range")
    if args.stddev is not None and args.stddev <= 0:
        raise ToolError("--stddev must be greater than 0")

    submissions = fetch_submissions_for_homework(
        args.cookie_file,
        course,
        clazz,
        str(work["work_id"]),
        3,
        args.timeout,
    )
    if args.limit is not None and args.limit <= 0:
        raise ToolError("--limit must be greater than 0")
    if args.limit is not None:
        submissions = submissions[: args.limit]
    if not submissions:
        raise ToolError(f"no pending submissions found for homework: {work['title']}")

    plan = []
    for row in submissions:
        plan.append(
            {
                "index": row.get("index"),
                "answer_id": row.get("answer_id"),
                "student_name": row.get("student_name"),
                "student_no": row.get("student_no"),
                "score": random_score(
                    low,
                    high,
                    args.decimals,
                    args.distribution,
                    args.mean,
                    args.stddev,
                ),
            }
        )

    distribution_summary = {
        "min": low,
        "max": high,
        "decimals": args.decimals,
        "distribution": args.distribution,
        "mean": (low + high) / 2 if args.mean is None else args.mean,
        "stddev": (high - low) / 6 if args.stddev is None else args.stddev,
    }
    results = []
    plan_payload = {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "course": course_summary(course),
        "selected_class": class_summary(clazz),
        "homework": work,
        "work_info": work_info,
        "score_range": distribution_summary,
        "plan": plan,
    }
    plan_file = ""
    if not args.commit:
        plan_file = save_grade_plan(plan_payload, args.plan_dir)

    if args.commit:
        for item in plan:
            result = mark_homework_score(
                session,
                course,
                clazz,
                str(work["work_id"]),
                str(item["answer_id"]),
                str(item["score"]),
                args.timeout,
            )
            results.append(
                {
                    "answer_id": item["answer_id"],
                    "student_name": item["student_name"],
                    "score": item["score"],
                    "status": bool(result.get("status")),
                    "message": result.get("msg", ""),
                    "response": result,
                }
            )

    if args.format == "json":
        print(
            json.dumps(
                {
                    "course": course_summary(course),
                    "selected_class": class_summary(clazz),
                    "homework": work,
                    "work_info": work_info,
                    "score_range": distribution_summary,
                    "dry_run": not args.commit,
                    "plan_file": plan_file,
                    "plan": plan,
                    "results": results,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(f"homework: {work['title']} (work_id={work['work_id']}, full_score={full_score:g})")
        print(f"score_range: {low:g}-{high:g}, decimals={args.decimals}")
        print(
            f"distribution: {args.distribution}"
            f" (mean={distribution_summary['mean']:g}, stddev={distribution_summary['stddev']:g})"
        )
        print(f"dry_run: {not args.commit}")
        if plan_file:
            print(f"plan_file: {plan_file}")
        print("final_score_summary:")
        for item in plan:
            print(
                f"  {item['index']}. {item['student_name']} "
                f"(answer_id={item['answer_id']}, student_no={item['student_no']}) -> {item['score']}"
            )
        if args.commit:
            print("results:")
            for result in results:
                print(
                    f"  {result['student_name']} answer_id={result['answer_id']} "
                    f"score={result['score']} status={result['status']} message={result['message']}"
                )
        else:
            print("not submitted; ask the user to confirm, then run homework-submit-plan with --commit")
    return 0


def command_homework_submit_plan(args: argparse.Namespace) -> int:
    if args.auto_login:
        ensure_api_logged_in(args)
    payload = load_grade_plan(args.plan_file)
    plan = payload.get("plan")
    course = payload.get("course")
    clazz = payload.get("selected_class")
    homework = payload.get("homework")
    if not isinstance(plan, list) or not isinstance(course, dict) or not isinstance(clazz, dict) or not isinstance(homework, dict):
        raise ToolError(f"invalid grade plan: {args.plan_file}")

    results = []
    verification = {}
    if args.commit:
        session = request_session_from_cookie_file(args.cookie_file)
        for item in plan:
            if not isinstance(item, dict):
                continue
            result = mark_homework_score(
                session,
                course,
                clazz,
                str(homework.get("work_id") or ""),
                str(item.get("answer_id") or ""),
                str(item.get("score") or ""),
                args.timeout,
            )
            results.append(
                {
                    "answer_id": item.get("answer_id"),
                    "student_name": item.get("student_name"),
                    "score": item.get("score"),
                    "status": bool(result.get("status")),
                    "message": result.get("msg", ""),
                    "response": result,
                }
            )
        if not args.skip_verify:
            verification = verify_submitted_scores(
                args.cookie_file,
                course,
                clazz,
                homework,
                [item for item in plan if isinstance(item, dict)],
                args.timeout,
            )

    if args.format == "json":
        print(
            json.dumps(
                {
                    "plan_file": args.plan_file,
                    "dry_run": not args.commit,
                    "plan": payload,
                    "results": results,
                    "verification": verification,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print(f"plan_file: {args.plan_file}")
        print_grade_plan_summary(payload)
        if args.commit:
            print("submitted: true")
            print("results:")
            for result in results:
                print(
                    f"  {result['student_name']} answer_id={result['answer_id']} "
                    f"score={result['score']} status={result['status']} message={result['message']}"
                )
            if verification:
                print(
                    "verification: "
                    f"expected={verification.get('expected')}, "
                    f"matched={verification.get('matched')}, "
                    f"mismatch={verification.get('mismatch')}, "
                    f"missing={verification.get('missing')}"
                )
                for row in verification.get("mismatches", []):
                    if not isinstance(row, dict):
                        continue
                    print(
                        f"  mismatch: {row.get('student_name')} answer_id={row.get('answer_id')} "
                        f"expected={row.get('expected')} actual={row.get('actual')}"
                    )
        else:
            print("not submitted; add --commit after the user confirms this exact summary")
    return 0


def command_login(args: argparse.Namespace) -> int:
    return api_login(args)


def command_browser_login(args: argparse.Namespace) -> int:
    ensure_agent_browser()
    browser_args = base_args(args.session, True)
    if args.fill:
        username, password = prompt_credentials()
        return fill_and_submit_password_login(browser_args, args.url, username, password)
    return manual_browser_login(browser_args, args.url, args.timeout, args.interval)


def command_password_login(args: argparse.Namespace) -> int:
    ensure_agent_browser()
    username = os.environ.get("CHAOXING_USERNAME", "").strip()
    password = os.environ.get("CHAOXING_PASSWORD", "")
    if not username or args.prompt:
        username, password = prompt_credentials()
    elif not password:
        password = getpass.getpass("Chaoxing password: ")
        if not password:
            raise ToolError("password cannot be empty")

    browser_args = base_args(args.session, True)
    return fill_and_submit_password_login(browser_args, args.url, username, password)


def command_open(args: argparse.Namespace) -> int:
    ensure_agent_browser()
    if args.auto_login:
        ensure_api_logged_in(args)
    browser_args = base_args(args.session, args.headed)
    injected = inject_cookie_file_into_browser(browser_args, args.cookie_file)
    if injected:
        print(f"injected_cookies: {injected}")
    run_browser([*browser_args, "open", args.url])
    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)
    return print_status(args.session, args.headed)


def command_status(args: argparse.Namespace) -> int:
    ensure_agent_browser()
    if args.auto_login:
        ensure_api_logged_in(args)
    browser_args = base_args(args.session, args.headed)
    inject_cookie_file_into_browser(browser_args, args.cookie_file)
    run_browser([*browser_args, "open", args.url])
    run_browser([*browser_args, "wait", "--load", "networkidle"], check=False)

    interval = max(args.interval, 0.25)
    deadline = time.time() + args.timeout
    last_code = 1
    while True:
        last_code = print_status(args.session, args.headed)
        if last_code == 0 or time.time() >= deadline:
            return last_code
        time.sleep(interval)


def command_clear(args: argparse.Namespace) -> int:
    ensure_agent_browser()
    browser_args = base_args(args.session, False)
    run_browser([*browser_args, "open", "about:blank"], check=False)
    run_browser([*browser_args, "cookies", "clear"], check=False)
    run_browser(["state", "clear", args.session], check=False)
    if os.path.exists(args.cookie_file):
        os.remove(args.cookie_file)
        print(f"removed cookie file: {args.cookie_file}")
    print(f"cleared cookies and saved browser state for session: {args.session}")
    return 0


def add_cookie_file_argument(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "--cookie-file",
        default=default_cookie_file(),
        help=f"saved API login cookie file, default: {default_cookie_file()}",
    )


def add_auto_login_argument(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "--no-auto-login",
        dest="auto_login",
        action="store_false",
        help="do not prompt for account/password when saved cookies are missing or expired",
    )
    parser.set_defaults(auto_login=True)


def add_login_check_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    parser.add_argument("--login-timeout", type=float, default=20.0, help="API login/check timeout in seconds")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Automate opening the SCUJCC Chaoxing/Fanya portal with a saved browser session."
    )
    parser.add_argument(
        "--session",
        default=DEFAULT_SESSION_NAME,
        help=f"agent-browser session name, default: {DEFAULT_SESSION_NAME}",
    )
    parser.add_argument("--url", default=PORTAL_URL, help=f"default: {PORTAL_URL}")

    subparsers = parser.add_subparsers(dest="command", required=True)

    setup = subparsers.add_parser(
        "setup",
        help="install Python dependencies into a local project directory",
    )
    setup.add_argument(
        "--target",
        default=default_local_deps_dir(),
        help=f"local dependency directory, default: {default_local_deps_dir()}",
    )
    setup.add_argument(
        "--requirements",
        default=default_requirements_file(),
        help=f"requirements file, default: {default_requirements_file()}",
    )
    setup.add_argument("--upgrade", action="store_true", help="upgrade packages in the local dependency directory")
    setup.set_defaults(func=command_setup)

    doctor = subparsers.add_parser(
        "doctor",
        help="check dependencies, optional browser support, cookie file, and saved login",
    )
    add_cookie_file_argument(doctor)
    doctor.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    doctor.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    doctor.add_argument(
        "--check-login",
        action="store_true",
        help="also verify that the saved cookie reaches an authenticated page",
    )
    doctor.set_defaults(func=command_doctor)

    login = subparsers.add_parser("login", help="login through Chaoxing HTTP APIs and save cookies")
    add_cookie_file_argument(login)
    login.add_argument("--prompt", action="store_true", help="always prompt, ignoring env vars")
    login.add_argument(
        "--macos-dialog",
        action="store_true",
        help="prompt for username/password using macOS secure dialogs instead of stdin",
    )
    login.add_argument(
        "--stdin-json",
        action="store_true",
        help="read username/password JSON from stdin; only use when the user explicitly accepts chat/input exposure",
    )
    login.add_argument("--fid", default="-1", help="Chaoxing school/org fid, default: -1")
    login.add_argument(
        "--timeout",
        type=float,
        default=20.0,
        help="request timeout in seconds",
    )
    login.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    login.set_defaults(func=command_login)

    browser_login = subparsers.add_parser(
        "browser-login",
        help="fallback: open a visible browser login window and save session cookies",
    )
    browser_login.add_argument(
        "--fill",
        action="store_true",
        help="prompt for username/password in terminal and submit them in the browser",
    )
    browser_login.add_argument(
        "--timeout",
        type=float,
        default=120.0,
        help="seconds to wait for browser login before asking for Enter",
    )
    browser_login.add_argument("--interval", type=float, default=2.0, help="poll interval in seconds")
    browser_login.add_argument("--manual", action="store_true", help=argparse.SUPPRESS)
    browser_login.set_defaults(func=command_browser_login)

    password_login = subparsers.add_parser(
        "password-login",
        help="fill username/password, prompting when env vars are absent",
    )
    password_login.add_argument("--prompt", action="store_true", help="always prompt, ignoring env vars")
    password_login.set_defaults(func=command_password_login)

    open_cmd = subparsers.add_parser("open", help="open the portal using the saved session")
    add_cookie_file_argument(open_cmd)
    add_auto_login_argument(open_cmd)
    add_login_check_arguments(open_cmd)
    open_cmd.add_argument("--headed", action="store_true", help="show the browser window")
    open_cmd.set_defaults(func=command_open)

    status = subparsers.add_parser("status", help="check whether the saved session reaches the portal")
    add_cookie_file_argument(status)
    add_auto_login_argument(status)
    add_login_check_arguments(status)
    status.add_argument("--headed", action="store_true", help="show the browser window")
    status.add_argument("--timeout", type=float, default=1.0, help="seconds to keep polling")
    status.add_argument("--interval", type=float, default=1.0, help="poll interval in seconds")
    status.set_defaults(func=command_status)

    api_status_cmd = subparsers.add_parser("api-status", help="check saved cookies with HTTP requests")
    add_cookie_file_argument(api_status_cmd)
    add_auto_login_argument(api_status_cmd)
    api_status_cmd.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    api_status_cmd.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    api_status_cmd.set_defaults(func=api_status)

    exam_question_draft = subparsers.add_parser(
        "exam-question-draft",
        aliases=["question-draft", "draft-exam-questions"],
        help="generate an offline exam-question draft from course content files",
    )
    exam_question_draft.add_argument("--format", choices=("table", "json", "markdown"), default="table")
    exam_question_draft.add_argument("--title", default="", help="exam/question draft title")
    exam_question_draft.add_argument("--course", default="", help="course name for metadata")
    exam_question_draft.add_argument("--chapter", default="", help="chapter/section name for metadata")
    exam_question_draft.add_argument("--knowledge-points", default="", help="comma-separated knowledge points")
    exam_question_draft.add_argument("--count", type=int, default=10, help="number of questions to draft")
    exam_question_draft.add_argument(
        "--types",
        default="single_choice,true_false,short_answer",
        help="question types, for example: single_choice,multiple_choice,true_false,short_answer,essay",
    )
    exam_question_draft.add_argument(
        "--difficulty",
        choices=("easy", "medium", "hard", "mixed"),
        default="mixed",
        help="question difficulty metadata, default: mixed",
    )
    exam_question_draft.add_argument(
        "--max-source-chars",
        type=int,
        default=24000,
        help="max characters to read from all source files",
    )
    exam_question_draft.add_argument(
        "--draft-dir",
        default=default_question_draft_dir(),
        help=f"directory for exam question drafts, default: {default_question_draft_dir()}",
    )
    exam_question_draft.add_argument("--output", default="", help="write JSON draft to this path")
    exam_question_draft.add_argument("--markdown-output", default="", help="also write a Markdown review copy")
    exam_question_draft.add_argument(
        "sources",
        nargs="+",
        help="source files: txt/md/html/docx/pdf/json, or '-' for stdin",
    )
    exam_question_draft.set_defaults(func=command_exam_question_draft)

    courses = subparsers.add_parser("courses", help="list teaching courses only")
    add_cookie_file_argument(courses)
    add_auto_login_argument(courses)
    courses.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    courses.add_argument("--login-timeout", type=float, default=20.0, help="API login/check timeout in seconds")
    courses.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    courses.add_argument("--format", choices=("table", "json"), default="table")
    courses.set_defaults(func=command_courses)

    course_classes = subparsers.add_parser(
        "course-classes",
        aliases=["enter-course"],
        help="show class list for a teaching course by name or course_id",
    )
    add_cookie_file_argument(course_classes)
    add_auto_login_argument(course_classes)
    course_classes.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    course_classes.add_argument("--login-timeout", type=float, default=20.0, help="API login/check timeout in seconds")
    course_classes.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    course_classes.add_argument("--format", choices=("table", "json"), default="table")
    course_classes.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    course_classes.add_argument("course", help="teaching course name, partial name, or course_id")
    course_classes.set_defaults(func=command_course_classes)

    task_options_cmd = subparsers.add_parser(
        "task-options",
        aliases=["class-tasks"],
        help="show exam/homework task choices for a selected teaching class",
    )
    add_cookie_file_argument(task_options_cmd)
    add_auto_login_argument(task_options_cmd)
    task_options_cmd.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    task_options_cmd.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    task_options_cmd.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    task_options_cmd.add_argument("--format", choices=("table", "json"), default="table")
    task_options_cmd.add_argument(
        "--owner-name",
        default="",
        help="override current user name used to prioritize classes",
    )
    task_options_cmd.add_argument("course", help="teaching course name, partial name, or course_id")
    task_options_cmd.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    task_options_cmd.set_defaults(func=command_task_options)

    course_task = subparsers.add_parser(
        "course-task",
        aliases=["enter-task"],
        help="generate or open the selected class's exam/homework interface",
    )
    add_cookie_file_argument(course_task)
    add_auto_login_argument(course_task)
    course_task.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    course_task.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    course_task.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    course_task.add_argument("--format", choices=("table", "json"), default="table")
    course_task.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    course_task.add_argument("--open", action="store_true", help="open the task interface in the browser session")
    course_task.add_argument("--headed", action="store_true", help="show the browser window when --open is used")
    course_task.add_argument("course", help="teaching course name, partial name, or course_id")
    course_task.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    course_task.add_argument("task", help="homework/作业 or exam/考试; exam creation/publishing is not automated")
    course_task.set_defaults(func=command_course_task)

    homework_ungraded = subparsers.add_parser(
        "homework-ungraded",
        aliases=["ungraded-homeworks"],
        help="list ungraded homework items for a selected teaching class",
    )
    add_cookie_file_argument(homework_ungraded)
    add_auto_login_argument(homework_ungraded)
    homework_ungraded.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_ungraded.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_ungraded.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_ungraded.add_argument("--format", choices=("table", "json"), default="table")
    homework_ungraded.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    homework_ungraded.add_argument("--all", action="store_true", help="show all homework items, not only pending ones")
    homework_ungraded.add_argument("course", help="teaching course name, partial name, or course_id")
    homework_ungraded.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    homework_ungraded.set_defaults(func=command_homework_ungraded)

    homework_submissions = subparsers.add_parser(
        "homework-submissions",
        aliases=["pending-submissions"],
        help="list pending submissions for a selected homework",
    )
    add_cookie_file_argument(homework_submissions)
    add_auto_login_argument(homework_submissions)
    homework_submissions.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_submissions.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_submissions.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_submissions.add_argument("--format", choices=("table", "json"), default="table")
    homework_submissions.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    homework_submissions.add_argument("--all-homeworks", action="store_true", help="resolve homework from all items instead of ungraded items")
    homework_submissions.add_argument("--status", type=int, default=3, help="mark status: 3 pending, 4 graded, 0 all")
    homework_submissions.add_argument("course", help="teaching course name, partial name, or course_id")
    homework_submissions.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    homework_submissions.add_argument("work", help="homework index, title, partial title, or work_id")
    homework_submissions.set_defaults(func=command_homework_submissions)

    homework_review = subparsers.add_parser(
        "homework-review",
        help="fetch or open a selected homework submission review page",
    )
    add_cookie_file_argument(homework_review)
    add_auto_login_argument(homework_review)
    homework_review.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_review.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_review.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_review.add_argument("--format", choices=("table", "json"), default="table")
    homework_review.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    homework_review.add_argument("--all-homeworks", action="store_true", help="resolve homework from all items instead of ungraded items")
    homework_review.add_argument("--status", type=int, default=3, help="mark status: 3 pending, 4 graded, 0 all")
    homework_review.add_argument("--max-chars", type=int, default=1200, help="max extracted review text characters")
    homework_review.add_argument("--open", action="store_true", help="open the review page in the browser session")
    homework_review.add_argument("--headed", action="store_true", help="show the browser window when --open is used")
    homework_review.add_argument("course", help="teaching course name, partial name, or course_id")
    homework_review.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    homework_review.add_argument("work", help="homework index, title, partial title, or work_id")
    homework_review.add_argument("answer", nargs="?", default="", help="submission index, answer_id, student name, or student number")
    homework_review.set_defaults(func=command_homework_review)

    homework_review_bundle = subparsers.add_parser(
        "homework-review-bundle",
        help="extract pending submission content for LLM-assisted grading",
    )
    add_cookie_file_argument(homework_review_bundle)
    add_auto_login_argument(homework_review_bundle)
    homework_review_bundle.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_review_bundle.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_review_bundle.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_review_bundle.add_argument("--format", choices=("table", "json"), default="table")
    homework_review_bundle.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    homework_review_bundle.add_argument("--all-homeworks", action="store_true", help="resolve homework from all items instead of ungraded items")
    homework_review_bundle.add_argument("--status", type=int, default=3, help="mark status: 3 pending, 4 graded, 0 all")
    homework_review_bundle.add_argument("--max-chars", type=int, default=1600, help="max extracted review text characters per student")
    homework_review_bundle.add_argument("--limit", type=int, default=None, help="limit number of submissions to extract")
    homework_review_bundle.add_argument("--download-assets", action="store_true", help="download answer images locally for model review")
    homework_review_bundle.add_argument(
        "--bundle-dir",
        default=default_review_bundle_dir(),
        help=f"directory for review bundles, default: {default_review_bundle_dir()}",
    )
    homework_review_bundle.add_argument("course", help="teaching course name, partial name, or course_id")
    homework_review_bundle.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    homework_review_bundle.add_argument("work", help="homework index, title, partial title, or work_id")
    homework_review_bundle.set_defaults(func=command_homework_review_bundle)

    homework_grade_reviewed = subparsers.add_parser(
        "homework-grade-reviewed",
        help="create a dry-run grade plan from a review bundle and LLM score reasons",
    )
    homework_grade_reviewed.add_argument("--format", choices=("table", "json"), default="table")
    homework_grade_reviewed.add_argument("--score-range", default="", help="allowed score interval, for example 70-95")
    homework_grade_reviewed.add_argument("--floor", type=float, default=70.0, help="scores below this require a concrete reason")
    homework_grade_reviewed.add_argument("--decimals", type=int, choices=(0, 1), default=0, help="score decimals")
    homework_grade_reviewed.add_argument("--allow-missing-reasons", action="store_true", help="do not require a reason for every score")
    homework_grade_reviewed.add_argument("--allow-partial", action="store_true", help="allow scoring only part of the review bundle")
    homework_grade_reviewed.add_argument(
        "--plan-dir",
        default=default_grade_plan_dir(),
        help=f"directory for dry-run grade plans, default: {default_grade_plan_dir()}",
    )
    homework_grade_reviewed.add_argument("review_bundle", help="JSON file from homework-review-bundle")
    homework_grade_reviewed.add_argument("scores_file", help="JSON with scores/reasons from content review")
    homework_grade_reviewed.set_defaults(func=command_homework_grade_reviewed)

    homework_grade_random = subparsers.add_parser(
        "homework-grade-random",
        help="randomly assign scores to pending homework submissions; dry-run by default",
    )
    add_cookie_file_argument(homework_grade_random)
    add_auto_login_argument(homework_grade_random)
    homework_grade_random.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_grade_random.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_grade_random.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_grade_random.add_argument("--format", choices=("table", "json"), default="table")
    homework_grade_random.add_argument("--owner-name", default="", help="override current user name used to prioritize classes")
    homework_grade_random.add_argument("--all-homeworks", action="store_true", help="resolve homework from all items instead of ungraded items")
    homework_grade_random.add_argument("--score-range", required=True, help="score interval, for example 80-90")
    homework_grade_random.add_argument("--decimals", type=int, choices=(0, 1), default=0, help="random score decimals")
    homework_grade_random.add_argument(
        "--distribution",
        choices=("uniform", "normal"),
        default="uniform",
        help="score distribution, default: uniform",
    )
    homework_grade_random.add_argument("--mean", type=float, default=None, help="normal distribution mean")
    homework_grade_random.add_argument("--stddev", type=float, default=None, help="normal distribution standard deviation")
    homework_grade_random.add_argument("--limit", type=int, default=None, help="limit number of pending submissions to grade")
    homework_grade_random.add_argument("--seed", type=int, default=None, help="random seed for reproducible dry runs")
    homework_grade_random.add_argument("--commit", action="store_true", help="actually submit scores to Chaoxing")
    homework_grade_random.add_argument(
        "--plan-dir",
        default=default_grade_plan_dir(),
        help=f"directory for dry-run grade plans, default: {default_grade_plan_dir()}",
    )
    homework_grade_random.add_argument("course", help="teaching course name, partial name, or course_id")
    homework_grade_random.add_argument("clazz", help="class index, class name, partial class name, or clazz_id")
    homework_grade_random.add_argument("work", help="homework index, title, partial title, or work_id")
    homework_grade_random.set_defaults(func=command_homework_grade_random)

    homework_submit_plan = subparsers.add_parser(
        "homework-submit-plan",
        aliases=["submit-grade-plan"],
        help="show or submit a saved homework grade plan after user confirmation",
    )
    add_cookie_file_argument(homework_submit_plan)
    add_auto_login_argument(homework_submit_plan)
    homework_submit_plan.add_argument("--check-url", default=AUTH_CHECK_URL, help=f"default: {AUTH_CHECK_URL}")
    homework_submit_plan.add_argument(
        "--login-timeout",
        type=float,
        default=20.0,
        help="API login/check timeout in seconds",
    )
    homework_submit_plan.add_argument("--timeout", type=float, default=20.0, help="request timeout in seconds")
    homework_submit_plan.add_argument("--format", choices=("table", "json"), default="table")
    homework_submit_plan.add_argument("--commit", action="store_true", help="actually submit the saved scores")
    homework_submit_plan.add_argument("--skip-verify", action="store_true", help="skip fetching the mark list after submit")
    homework_submit_plan.add_argument("plan_file", help="saved grade plan JSON from homework-grade-random")
    homework_submit_plan.set_defaults(func=command_homework_submit_plan)

    clear = subparsers.add_parser("clear", help="clear cookies from this automation session")
    add_cookie_file_argument(clear)
    clear.set_defaults(func=command_clear)

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        return args.func(args)
    except KeyboardInterrupt:
        print("\ninterrupted", file=sys.stderr)
        return 130
    except ToolError as exc:
        print(str(exc), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
